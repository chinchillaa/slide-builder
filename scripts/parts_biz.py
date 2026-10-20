# -*- coding: utf-8 -*-
"""parts_biz - 業務資料で頻出の型（会議・報告、リスク・課題、費用・効果、要件、目標、顧客、意思決定、フレームワーク）"""
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from pptx_helpers import *  # noqa: F401,F403
from pptx_helpers import (_STATE, _arrow_line, _blank, _box, _chrome_footer, _circle, _line, _put_text, _rect,
                          _rgb, _split, _textbox, chart, estimate_height, fit_one_line, part, text_width)
from parts_connect import connect, icon, icon_label
from parts_data import progress_bars
from parts_flow import _chevrons
from parts_structure import logic_tree
from parts_table import _STATUS, _native_table
from parts_text import grid_areas, split_cols, split_rows, tag, titled_box

_SIGNAL = {"良好": COLOR_STATUS_OK, "注意": COLOR_STATUS_WARN, "問題": COLOR_STATUS_NG,
           "green": COLOR_STATUS_OK, "yellow": COLOR_STATUS_WARN, "red": COLOR_STATUS_NG}
_LEVEL = {"高": COLOR_ACCENT, "中": COLOR_PRIMARY_MID, "低": COLOR_GRAY_DARK}


def _hdr_style(i, j, c):
    return COLOR_PRIMARY, COLOR_WHITE, True, "center"


# ---------------------------------------------------------------- 会議・報告
@part("会議・報告", "信号機式の状況報告（項目ごとに 良好/注意/問題 の信号とコメント）")
def traffic_light(slide, area, rows, size=12):
    """rows: [(項目, 状態 "良好"|"注意"|"問題", コメント)]"""
    x, y, w, h = area
    cols = split_cols((x, y, w, 26), [1.2, 0.95, 3.1], 4)
    for (cx, cy, cw, ch), t in zip(cols, ["項目", "状況", "コメント"]):
        _box(slide, cx, cy, cw, ch, t, style="label", size=size)
    for i, ((item, st, cm), (rx, ry, rw, rh)) in enumerate(zip(rows, split_rows((x, y + 30, w, h - 30), [1] * len(rows), 4))):
        c0, c1, c2 = [(c[0], c[2]) for c in cols]
        _box(slide, c0[0], ry, c0[1], rh, item, style="gray", size=size, where=f"traffic_light項目[{i}]")
        _rect(slide, c1[0], ry, c1[1], rh, fill=COLOR_WHITE, line=COLOR_GRAY_DARK)
        d = min(22, rh - 10)
        _circle(slide, c1[0] + 10, ry + (rh - d) / 2, d, None, fill=_SIGNAL[st])
        if text_width(st, size) * 1.1 > c1[1] - d - 22:
            raise ValueError(f"[traffic_light] 状況欄が狭く「{st}」が1行に収まりません。領域の幅を広げてください。")
        _textbox(slide, c1[0] + 14 + d, ry, c1[1] - d - 16, rh, st, size=size, bold=True, anchor="middle",
                 color=_SIGNAL[st] if st != "注意" else COLOR_TEXT, where="traffic_light状態")
        _box(slide, c2[0], ry, c2[1], rh, cm, style="text", size=size, where=f"traffic_lightコメント[{i}]")


_DTYPE = {"決定": "primary", "宿題": "mid", "確認": "gray", "保留": "outline"}


@part("会議・報告", "決定事項・宿題・確認事項の一覧（種別タグ＋内容＋担当・期限）")
def decision_log(slide, area, items, size=12):
    """items: [(種別 "決定"|"宿題"|"確認"|"保留", 内容, 担当 or "", 期限 or "")]"""
    for i, ((kind, body, who, due), (x, y, w, h)) in enumerate(zip(items, split_rows(area, [1] * len(items), 4))):
        tag(slide, x, y + h / 2 - 10, kind, style=_DTYPE.get(kind, "gray"))
        meta = " / ".join(t for t in (f"担当：{who}" if who else "", f"期限：{due}" if due else "") if t)
        mw = 170 if meta else 0
        _textbox(slide, x + 52, y, w - 52 - mw, h, body, size=size, anchor="middle", where=f"decision_log[{i}]")
        if meta:
            _textbox(slide, x + w - mw, y, mw, h, meta, size=11, color=COLOR_TEXT_SUB, align="right", anchor="middle",
                     dense=True, where="decision_log担当")
        _line(slide, x, y + h + 2, x + w, y + h + 2, COLOR_GRAY_LIGHT, LINE_W)


@part("会議・報告", "会議体一覧（会議名・頻度・参加者・目的）")
def meeting_bodies(slide, area, rows, size=12, stretch=False):
    data = [["会議体", "頻度", "参加者", "目的・議題"]] + [list(r) for r in rows]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        if j == 0:
            return COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, "left"
        return COLOR_WHITE, COLOR_TEXT, False, "center" if j == 1 else "left"
    return _native_table(slide, area, data, [1.3, 0.9, 2, 3], st, size, stretch=stretch, where="meeting_bodies")


@part("会議・報告", "週次報告（全体状況の信号＋今週の実績／来週の予定／課題・リスク）")
def weekly_report(slide, area, overall, done, plan, issues, period=None, size=12):
    """overall = (状態 "良好"|"注意"|"問題", コメント)"""
    x, y, w, h = area
    st, cm = overall
    _box(slide, x, y, 110, 40, "全体状況", style="label", size=size + 1)
    _rect(slide, x + 110, y, w - 110, 40, fill=COLOR_WHITE, line=COLOR_GRAY_DARK)
    _circle(slide, x + 122, y + 9, 22, None, fill=_SIGNAL[st])
    pw = 200 if period else 0
    _textbox(slide, x + 152, y, w - 160 - pw, 40, f"**{st}**　{cm}", size=size + 1, anchor="middle", where="weekly_report状況")
    if period:
        _textbox(slide, x + w - pw - 8, y, pw, 40, period, size=11, color=COLOR_TEXT_SUB, align="right", anchor="middle",
                 dense=True, where="weekly_report期間")
    for a, (t, items, s) in zip(split_cols((x, y + 52, w, h - 52), [1, 1, 1]),
                                (("今週の実績", done, "primary"), ("来週の予定", plan, "mid"), ("課題・リスク", issues, "gray"))):
        titled_box(slide, a, t, items, style=s, size=size)


@part("会議・報告", "会議のアジェンダ（時間・議題・担当）")
def meeting_agenda(slide, area, items, size=13):
    """items: [(時間, 議題, 担当)]"""
    for i, ((tm, topic, who), (x, y, w, h)) in enumerate(zip(items, split_rows(area, [1] * len(items), 4))):
        _box(slide, x, y, 130, h, tm, style="pale", size=size, where=f"meeting_agenda時間[{i}]")
        _textbox(slide, x + 142, y, w - 142 - 150, h, topic, size=size, bold=True, anchor="middle",
                 where=f"meeting_agenda議題[{i}]")
        _textbox(slide, x + w - 150, y, 150, h, who, size=12, color=COLOR_TEXT_SUB, align="right", anchor="middle",
                 where="meeting_agenda担当")
        _line(slide, x + 130, y + h + 2, x + w, y + h + 2, COLOR_GRAY_LIGHT, LINE_W)


@part("会議・報告", "変更履歴（版・日付・変更内容・担当）")
def revision_history(slide, area, rows, size=12, stretch=False):
    data = [["版", "日付", "変更内容", "担当"]] + [list(r) for r in rows]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        return COLOR_WHITE, COLOR_TEXT, j == 0, "center" if j < 2 else "left"
    return _native_table(slide, area, data, [0.6, 1.2, 5, 1.2], st, size, stretch=stretch, where="revision_history")


@part("会議・報告", "お知らせ・注意の枠（info / warn）")
def notice_box(slide, area, title, items, kind="info", size=13):
    x, y, w, h = area
    col = COLOR_PRIMARY_MID if kind == "info" else COLOR_STATUS_WARN
    _rect(slide, x, y, w, h, fill=COLOR_WHITE, line=col, line_w=2)
    _rect(slide, x, y, 6, h, fill=col)
    if kind == "info":
        _circle(slide, x + 18, y + 12, 28, "i", fill=col, size=16)
    else:
        icon(slide, x + 18, y + 12, 28, "warning")
    _textbox(slide, x + 56, y + 10, w - 66, 30, title, size=size + 1, bold=True,
             color=COLOR_PRIMARY if kind == "info" else COLOR_TEXT, anchor="middle", where="notice_box見出し")
    _textbox(slide, x + 56, y + 44, w - 66, h - 52, items, size=size, bullets=isinstance(items, list),
             where="notice_box本文")


# ---------------------------------------------------------------- リスク・課題
@part("リスク・課題", "リスクマトリクス（影響度×発生可能性の3×3に番号を配置＋右に一覧）")
def risk_matrix(slide, area, risks, size=12):
    """risks: [(リスク名, 影響度 1-3, 発生可能性 1-3)]"""
    x, y, w, h = area
    gw = min(w * 0.5, h - 30)
    gx, gy = x + 34, y
    cell = (gw - 4) / 3
    for r in range(3):
        for c in range(3):
            score = (3 - r) * (c + 1)
            fill = COLOR_PRIMARY_MID if score >= 6 else COLOR_PRIMARY_LIGHT if score >= 3 else COLOR_PRIMARY_PALE
            _rect(slide, gx + c * (cell + 2), gy + r * (cell + 2), cell, cell, fill=fill)
    _arrow_line(slide, gx - 6, gy + gw, gx - 6, gy, COLOR_PRIMARY, 1.5)
    _arrow_line(slide, gx, gy + gw + 6, gx + gw, gy + gw + 6, COLOR_PRIMARY, 1.5)
    tb = _textbox(slide, x, gy, 24, gw, "影響度", size=11, bold=True, color=COLOR_PRIMARY, align="center",
                  anchor="middle", dense=True, check=False)
    tb.text_frame._txBody.bodyPr.set("vert", "eaVert")
    _textbox(slide, gx, gy + gw + 10, gw, 20, "発生可能性", size=11, bold=True, color=COLOR_PRIMARY,
             align="center", dense=True, where="risk_matrix軸")
    placed = {}
    for i, (name, imp, lik) in enumerate(risks):
        k = placed.get((imp, lik), 0)
        placed[(imp, lik)] = k + 1
        cx = gx + (lik - 1) * (cell + 2) + 8 + (k % 3) * 28
        cy = gy + (3 - imp) * (cell + 2) + 8 + (k // 3) * 28
        _circle(slide, cx, cy, 24, i + 1, fill=COLOR_WHITE, color=COLOR_PRIMARY, size=12, line=COLOR_PRIMARY)
    lx = gx + gw + 30
    for i, ((name, imp, lik), (rx, ry, rw, rh)) in enumerate(zip(risks, split_rows((lx, y, x + w - lx, gw), [1] * len(risks), 4))):
        score = imp * lik
        lv = "高" if score >= 6 else "中" if score >= 3 else "低"
        _circle(slide, rx, ry + (rh - 22) / 2, 22, i + 1, fill=COLOR_PRIMARY, size=11)
        _textbox(slide, rx + 30, ry, rw - 80, rh, name, size=size, anchor="middle", where=f"risk_matrix一覧[{i}]")
        tag(slide, rx + rw - 44, ry + rh / 2 - 10, lv, style={"高": "alert", "中": "mid", "低": "gray"}[lv])


@part("リスク・課題", "リスク対応表（影響度・発生可能性を高中低で色分け）")
def risk_table(slide, area, rows, size=12, stretch=False):
    """rows: [(リスク, 影響度 高中低, 発生可能性 高中低, 対策, 担当)]"""
    data = [["リスク", "影響度", "発生可能性", "対策", "担当"]] + [list(r) for r in rows]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        if j in (1, 2):
            return COLOR_WHITE, _LEVEL.get(c, COLOR_TEXT), True, "center"
        return COLOR_WHITE, COLOR_TEXT, j == 0, "left"
    return _native_table(slide, area, data, [2.2, 0.8, 1, 3, 1], st, size, stretch=stretch, where="risk_table")


@part("リスク・課題", "課題管理表（優先度を色分け、状況をステータス色で）")
def issue_tracker(slide, area, rows, size=12, stretch=False):
    """rows: [(課題, 優先度 高中低, 担当, 期限, 状況)]"""
    data = [["No", "課題", "優先度", "担当", "期限", "状況"]] + [[str(i + 1)] + list(r) for i, r in enumerate(rows)]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        if j == 2:
            return COLOR_WHITE, _LEVEL.get(c, COLOR_TEXT), True, "center"
        if j == 5:
            f, col, b = _STATUS.get(c, (COLOR_WHITE, COLOR_TEXT, False))
            return f, col, b, "center"
        return COLOR_WHITE, COLOR_TEXT, False, "center" if j in (0, 3, 4) else "left"
    return _native_table(slide, area, data, [0.6, 4, 0.9, 1.1, 1, 1], st, size, stretch=stretch, where="issue_tracker")


# ---------------------------------------------------------------- 費用・効果
def _num(v):
    if isinstance(v, (int, float)):
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.1f}"
    return v


@part("費用・効果", "費用内訳表（小計・合計行、金額は右寄せ・カンマ区切り）")
def cost_table(slide, area, header, rows, unit=None, auto_total=True, size=12):
    """rows: [(費目, 金額, ...)]。費目が「小計」で始まる行は小計として強調。auto_total で最終行に合計を追加"""
    x, y, w, h = area
    if unit:
        _textbox(slide, x + w - 200, y, 200, 18, f"（単位：{unit}）", size=10, color=COLOR_TEXT_SUB, align="right",
                 dense=True, where="cost_table単位")
        y, h = y + 20, h - 20
    rows = [list(r) for r in rows]
    if auto_total:
        tot = ["合計"] + [sum(r[j] for r in rows if not str(r[0]).startswith("小計") and isinstance(r[j], (int, float)))
                         for j in range(1, len(header))]
        rows.append(tot)
    data = [list(header)] + [[_num(c) for c in r] for r in rows]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        lab = str(data[i][0])
        al = "left" if j == 0 else "right"
        if lab == "合計":
            return COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, al
        if lab.startswith("小計"):
            return COLOR_GRAY_PALE, COLOR_TEXT, True, al
        return COLOR_WHITE, COLOR_TEXT, False, al
    return _native_table(slide, (x, y, w, h), data, [2] + [1] * (len(header) - 1), st, size, where="cost_table")


@part("費用・効果", "投資回収グラフ（累計投資と累計効果の折れ線＋回収時期）")
def roi_chart(slide, area, periods, investment, effect, unit=None, cumulative=True, size=12):
    """investment/effect: 期ごとの金額。cumulative=True なら累計に変換して描く"""
    x, y, w, h = area
    inv, eff = list(investment), list(effect)
    if cumulative:
        for k in range(1, len(inv)):
            inv[k] += inv[k - 1]
            eff[k] += eff[k - 1]
    pay = next((periods[k] for k in range(len(periods)) if eff[k] >= inv[k]), None)
    chart(slide, area, "line", periods, {"累計投資": inv, "累計効果": eff}, title="投資と効果の累計", unit=unit,
          labels=False, size=size)
    if pay:
        tag(slide, x + 6, y + 2, f"投資回収：{pay}", style="alert", size=12, h=22)


@part("費用・効果", "予算実績対比表（計画・実績・差異・達成率。未達は赤）")
def budget_actual_table(slide, area, rows, unit=None, size=12):
    """rows: [(項目, 計画, 実績)]"""
    x, y, w, h = area
    if unit:
        _textbox(slide, x + w - 200, y, 200, 18, f"（単位：{unit}）", size=10, color=COLOR_TEXT_SUB, align="right",
                 dense=True, where="budget単位")
        y, h = y + 20, h - 20
    body = [(n, p, a, a - p, a / p * 100 if p else 0) for n, p, a in rows]
    tp, ta = sum(r[1] for r in rows), sum(r[2] for r in rows)
    body.append(("合計", tp, ta, ta - tp, ta / tp * 100 if tp else 0))
    data = [["項目", "計画", "実績", "差異", "達成率"]] + [[n, _num(p), _num(a), f"{d:+,.0f}", f"{r:.0f}%"]
                                                     for n, p, a, d, r in body]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        tot = i == len(data) - 1
        fill = COLOR_PRIMARY_PALE if tot else COLOR_WHITE
        if j == 0:
            return fill, COLOR_PRIMARY if tot else COLOR_TEXT, tot, "left"
        bad = (j == 3 and body[i - 1][3] < 0) or (j == 4 and body[i - 1][4] < 100)
        return fill, COLOR_ACCENT if bad else (COLOR_PRIMARY if tot else COLOR_TEXT), bad or tot, "right"
    return _native_table(slide, (x, y, w, h), data, [2, 1, 1, 1, 1], st, size, where="budget_actual_table")


@part("費用・効果", "効果の内訳（定量効果の金額一覧＋合計／定性効果の箇条書き）")
def effect_breakdown(slide, area, quantitative, qualitative, unit="万円/年", size=12):
    """quantitative: [(項目, 金額)]"""
    a, b = split_cols(area, [1.1, 1], 20)
    x, y, w, h = a
    _box(slide, x, y, w, 30, f"定量効果（{unit}）", style="label", size=size + 1)
    _rect(slide, x, y + 30, w, h - 30, fill=COLOR_WHITE, line=COLOR_PRIMARY)
    rows = split_rows((x + 12, y + 40, w - 24, h - 92), [1] * len(quantitative), 2)
    for i, ((item, v), (rx, ry, rw, rh)) in enumerate(zip(quantitative, rows)):
        _textbox(slide, rx, ry, rw - 110, rh, item, size=size, anchor="middle", where=f"effect_breakdown[{i}]")
        _textbox(slide, rx + rw - 110, ry, 110, rh, _num(v), size=size, align="right", anchor="middle",
                 where="effect_breakdown金額")
        _line(slide, rx, ry + rh + 1, rx + rw, ry + rh + 1, COLOR_GRAY_LIGHT, LINE_W)
    tot = sum(v for _, v in quantitative)
    _line(slide, x + 12, y + h - 48, x + w - 12, y + h - 48, COLOR_PRIMARY, 1.5)
    _textbox(slide, x + 12, y + h - 44, w - 180, 36, "合計", size=size + 2, bold=True, color=COLOR_PRIMARY,
             anchor="middle", where="effect_breakdown合計")
    _textbox(slide, x + w - 160, y + h - 44, 148, 36, _num(tot), size=fit_one_line(_num(tot), 140, size + 6), bold=True,
             color=COLOR_PRIMARY, align="right", anchor="middle", where="effect_breakdown合計額")
    titled_box(slide, b, "定性効果", qualitative, style="mid", size=size)


# ---------------------------------------------------------------- 要件・仕様
_MOSCOW = {"Must": (COLOR_PRIMARY, COLOR_WHITE), "Should": (COLOR_PRIMARY_MID, COLOR_WHITE),
           "Could": (COLOR_PRIMARY_PALE, COLOR_PRIMARY), "Won't": (COLOR_GRAY_PALE, COLOR_TEXT_SUB)}


@part("要件・仕様", "要件一覧（ID・区分・要件・優先度 MoSCoW を色分け）")
def requirements_table(slide, area, rows, size=11, stretch=False):
    """rows: [(ID, 区分, 要件, 優先度 "Must"|"Should"|"Could"|"Won't")]"""
    data = [["ID", "区分", "要件", "優先度"]] + [list(r) for r in rows]

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        if j == 3:
            f, col = _MOSCOW.get(c, (COLOR_WHITE, COLOR_TEXT))
            return f, col, True, "center"
        return COLOR_WHITE, COLOR_TEXT, j == 0, "center" if j == 0 else "left"
    return _native_table(slide, area, data, [0.8, 1.2, 5, 1], st, size, stretch=stretch, where="requirements_table")


@part("要件・仕様", "グループ化した表（1列目の同じ値を縦に結合。機能一覧・非機能要件など）")
def grouped_table(slide, area, header, rows, col_ratios=None, size=12, stretch=False):
    rows = [list(r) for r in rows]
    spans, start = [], 0
    for i in range(1, len(rows) + 1):
        if i == len(rows) or rows[i][0] != rows[start][0]:
            spans.append((start, i - 1))
            start = i
    data = [list(header)]
    for i, r in enumerate(rows):
        data.append(([r[0]] if any(s == i for s, _ in spans) else [""]) + r[1:])

    def st(i, j, c):
        if i == 0:
            return _hdr_style(i, j, c)
        if j == 0:
            return COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, "center"
        return COLOR_WHITE, COLOR_TEXT, False, "left"
    gf = _native_table(slide, area, data, col_ratios or [1.2] + [2] * (len(header) - 1), st, size, stretch=stretch,
                       where="grouped_table")
    for s, e in spans:
        if e > s:
            gf.table.cell(s + 1, 0).merge(gf.table.cell(e + 1, 0))
    return gf


# ---------------------------------------------------------------- 目標管理
@part("目標管理", "KGI/KPI ツリー（各ノードに 現状 → 目標）")
def kpi_tree(slide, area, root, size=11):
    """root = (ラベル, 現状, 目標, [子, ...])"""
    def conv(n):
        lab, cur, tgt, kids = n if len(n) == 4 else (*n, [])
        t = f"**{lab}**\n{cur} → **{tgt}**"
        return (t, [conv(k) for k in kids]) if kids else t
    logic_tree(slide, area, conv(root), size=size)


@part("目標管理", "OKR（Objective と Key Results の進捗）")
def okr(slide, area, objective, krs, size=13):
    """krs: [(KR の文, 進捗 0-100)]"""
    x, y, w, h = area
    _box(slide, x, y, 130, 64, "Objective", style="label", size=size + 1)
    sp = _rect(slide, x + 130, y, w - 130, 64, fill=COLOR_PRIMARY_PALE)
    _put_text(sp, objective, size=size + 3, bold=True, color=COLOR_PRIMARY, anchor="middle", inset=(14, 4),
              where="okr目的")
    _textbox(slide, x, y + 76, 200, 22, "Key Results", size=size, bold=True, color=COLOR_PRIMARY_MID, where="okr見出し")
    progress_bars(slide, (x, y + 104, w, h - 104), [(f"KR{i + 1}　{t}", p) for i, (t, p) in enumerate(krs)],
                  label_ratio=0.45, size=size - 1)


@part("目標管理", "目標と実績の対比（目標＝枠、実績＝塗りの横棒と達成率）")
def target_actual(slide, area, items, unit="", size=12):
    """items: [(ラベル, 目標, 実績)]"""
    x, y, w, h = area
    vmax = max(max(t, a) for _, t, a in items) or 1
    lw, rw = w * 0.22, 130
    bw = w - lw - rw - 10
    from parts_text import legend
    legend(slide, x + w - 170, y, [("目標", COLOR_PRIMARY_PALE), ("実績", COLOR_PRIMARY_MID)])
    for i, ((lab, t, a), (rx, ry, _, rh)) in enumerate(zip(items, split_rows((x, y + 22, w, h - 22), [1] * len(items), 6))):
        bh = min(rh * 0.6, 26)
        by = ry + (rh - bh) / 2
        _textbox(slide, rx, ry, lw - 8, rh, lab, size=size, bold=True, anchor="middle", where=f"target_actual[{i}]")
        _rect(slide, rx + lw, by, bw * t / vmax, bh, fill=COLOR_PRIMARY_PALE, line=COLOR_PRIMARY_LIGHT)
        _rect(slide, rx + lw, by + bh * 0.2, bw * a / vmax, bh * 0.6, fill=COLOR_PRIMARY_MID)
        pct = a / t * 100 if t else 0
        _textbox(slide, rx + lw + bw + 10, ry, rw, rh, f"{_num(a)}/{_num(t)}{unit}　**{pct:.0f}%**", size=size,
                 color=COLOR_PRIMARY if pct >= 100 else COLOR_ACCENT, anchor="middle", dense=True,
                 where="target_actual値")


# ---------------------------------------------------------------- 顧客・利用者
@part("顧客・利用者", "カスタマージャーニー（段階×行動・接点・課題＋感情の曲線）")
def customer_journey(slide, area, stages, rows, emotions=None, label_w=90, size=11):
    """rows: [(行ラベル, [各段階のテキスト])]  emotions: 各段階の感情 -2〜+2"""
    x, y, w, h = area
    gx, gw = x + label_w + 6, w - label_w - 6
    cols = _chevrons(slide, gx, y, gw, 32, stages, size=size + 1)
    n_rows = len(rows) + (1 if emotions else 0)
    rhs = _split(h - 40, [1] * n_rows, 4)
    ry = y + 40
    for i, (lab, cells) in enumerate(rows):
        _box(slide, x, ry, label_w, rhs[i], lab, style="label", size=size + 1)
        for j, (cx, cw) in enumerate(cols):
            _box(slide, cx, ry, cw, rhs[i], cells[j], style="text", size=size, dense=True,
                 where=f"customer_journey[{i},{j}]", inset=(4, 2))
        ry += rhs[i] + 4
    if emotions:
        eh = rhs[-1]
        _box(slide, x, ry, label_w, eh, "感情", style="label", size=size + 1)
        _rect(slide, gx, ry, gw, eh, fill=COLOR_GRAY_PALE)
        mid = ry + eh / 2
        _line(slide, gx, mid, gx + gw, mid, COLOR_GRAY, 1, dash=True)
        pts = [(cx + cw / 2, mid - e / 2 * (eh / 2 - 10)) for (cx, cw), e in zip(cols, emotions)]
        fb = slide.shapes.build_freeform(Pt(pts[0][0]), Pt(pts[0][1]), scale=1.0)
        fb.add_line_segments([(Pt(px), Pt(py)) for px, py in pts[1:]], close=False)
        ln = fb.convert_to_shape()
        ln.fill.background()
        ln.line.color.rgb = _rgb(COLOR_PRIMARY_MID)
        ln.line.width = Pt(2.25)
        ln.shadow.inherit = False
        for (px, py), e in zip(pts, emotions):
            _circle(slide, px - 6, py - 6, 12, None, fill=COLOR_PRIMARY if e >= 0 else COLOR_ACCENT)


_LIKERT = ("そう思う", "ややそう思う", "どちらでもない", "あまりそう思わない", "そう思わない")


@part("顧客・利用者", "アンケート結果（5段階評価の100%帯グラフ。ネイティブ）")
def likert_chart(slide, area, questions, data, labels=_LIKERT, size=11):
    """data[設問] = [各段階の割合(%)]"""
    series = {lab: [row[k] for row in data] for k, lab in enumerate(labels)}
    gf = chart(slide, area, "bar_100", questions, series, number_format="0", size=size)
    ch = gf.chart
    ch.category_axis.reverse_order = True
    cols = [COLOR_PRIMARY, COLOR_PRIMARY_MID, COLOR_GRAY_LIGHT, COLOR_GRAY, COLOR_GRAY_DARK]
    for s, col in zip(ch.plots[0].series, cols):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = _rgb(col)
        dl = s.data_labels
        dl.show_value = True
        dl.number_format, dl.number_format_is_linked = "0", False
        dl.font.size = Pt(size)
        s.data_labels.font.color.rgb = _rgb(COLOR_TEXT if col in (COLOR_GRAY_LIGHT, COLOR_GRAY) else COLOR_WHITE)
    return gf


@part("顧客・利用者", "共感マップ（言う・考える・行動・感じる＋中央に人物）")
def empathy_map(slide, area, persona, says, thinks, does, feels, size=12):
    x, y, w, h = area
    for a, (t, items, st) in zip(grid_areas(area, 2, 2, 12), (("言っていること", says, "primary"),
                                                              ("考えていること", thinks, "mid"),
                                                              ("行動していること", does, "mid"),
                                                              ("感じていること", feels, "primary"))):
        titled_box(slide, a, t, items, style=st, size=size)
    d = min(80, h * 0.24)   # 見出し帯の角だけに重なる大きさ（本文は隠さない）
    sp = _circle(slide, x + w / 2 - d / 2, y + h / 2 - d / 2, d, None, fill=COLOR_PRIMARY_PALE, line=COLOR_PRIMARY)
    _put_text(sp, persona, size=size, bold=True, color=COLOR_PRIMARY, align="center", anchor="middle", inset=(2, 2),
              where="empathy_map人物")


# ---------------------------------------------------------------- 意思決定
@part("意思決定", "推奨案（選択肢カードで推奨を明示＋推奨理由・次のアクション）")
def recommendation(slide, area, options, recommended, reasons, actions, size=12):
    """options: [(案の名前, 概要)]"""
    x, y, w, h = area
    top, bot = split_rows(area, [0.8, 1], 14)
    for i, ((name, summ), (cx, cy, cw, ch)) in enumerate(zip(options, split_cols(top, [1] * len(options)))):
        on = i == recommended
        _rect(slide, cx, cy + 10, cw, ch - 10, fill=COLOR_PRIMARY_PALE if on else COLOR_WHITE,
              line=COLOR_PRIMARY if on else COLOR_GRAY, line_w=2.25 if on else LINE_W)
        if on:
            tag(slide, cx + 10, cy, "推奨", style="alert", size=11, h=20)
        _textbox(slide, cx + 10, cy + 24, cw - 20, 26, name, size=size + 2, bold=True,
                 color=COLOR_PRIMARY if on else COLOR_TEXT_SUB, align="center", where=f"recommendation案[{i}]")
        _textbox(slide, cx + 10, cy + 52, cw - 20, ch - 58, summ, size=size, align="center",
                 color=COLOR_TEXT if on else COLOR_TEXT_SUB, where=f"recommendation概要[{i}]")
    a, b = split_cols(bot, [1, 1], 16)
    titled_box(slide, a, "推奨する理由", reasons, style="primary", size=size)
    titled_box(slide, b, "次のアクション", actions, style="mid", size=size)


@part("意思決定", "前提条件と制約事項（番号付き）")
def assumptions(slide, area, premises, constraints, size=13):
    a, b = split_cols(area, [1, 1], 16)
    titled_box(slide, a, "前提条件", [f"**前提{i + 1}**　{t}" for i, t in enumerate(premises)], style="primary", size=size)
    titled_box(slide, b, "制約事項", [f"**制約{i + 1}**　{t}" for i, t in enumerate(constraints)], style="gray", size=size)


@part("意思決定", "料金・プラン比較カード（おすすめを強調）")
def plan_cards(slide, area, plans, recommended=None, size=12):
    """plans: [(プラン名, 価格表示, [特徴, ...])]"""
    for i, ((name, price, feats), (x, y, w, h)) in enumerate(zip(plans, split_cols(area, [1] * len(plans), 14))):
        on = i == recommended
        _rect(slide, x, y + 10, w, h - 10, fill=COLOR_WHITE, line=COLOR_PRIMARY_MID if on else COLOR_GRAY,
              line_w=2.25 if on else LINE_W)
        _box(slide, x, y + 10, w, 34, name, style="label2" if on else "label", size=size + 2)
        if on:
            tag(slide, x + w - 70, y, "おすすめ", style="alert", size=10, h=20)
        _textbox(slide, x + 8, y + 50, w - 16, 44, price, size=fit_one_line(price, w - 24, 24), bold=True,
                 color=COLOR_PRIMARY, align="center", anchor="middle", where=f"plan_cards価格[{i}]")
        _line(slide, x + 14, y + 98, x + w - 14, y + 98, COLOR_GRAY_LIGHT, LINE_W)
        _textbox(slide, x + 10, y + 104, w - 20, h - 112, feats, size=size, bullets=True, where=f"plan_cards特徴[{i}]")


# ---------------------------------------------------------------- 時系列
@part("スケジュール", "年表（期間×分野のマスに出来事を記入）")
def chronology(slide, area, periods, rows, events, label_w=110, size=11):
    """events: {(行index, 期間index): テキスト or list}"""
    x, y, w, h = area
    gx, gw = x + label_w + 4, w - label_w - 4
    pw = gw / len(periods)
    for j, p in enumerate(periods):
        _box(slide, gx + j * pw + 1, y, pw - 2, 26, p, style="pale", size=size + 1)
    rh = (h - 30) / len(rows)
    for i, lab in enumerate(rows):
        ry = y + 30 + i * rh
        _box(slide, x, ry + 1, label_w, rh - 2, lab, style="label", size=size + 1)
        if i:
            _line(slide, gx, ry, gx + gw, ry, COLOR_GRAY_LIGHT, LINE_W)
        for j in range(len(periods)):
            ev = events.get((i, j))
            if ev:
                _textbox(slide, gx + j * pw + 4, ry + 4, pw - 8, rh - 8, ev, size=size, bullets=isinstance(ev, list),
                         anchor="middle", dense=True, where=f"chronology[{i},{j}]")
    for j in range(1, len(periods)):
        _line(slide, gx + j * pw, y + 30, gx + j * pw, y + h, COLOR_GRAY_LIGHT, 0.5, dash=True)


# ---------------------------------------------------------------- フレームワーク
_FW = {"3c": ([("Customer：市場・顧客", "primary"), ("Competitor：競合", "mid"), ("Company：自社", "pale")], (1, 3)),
       "4p": ([("Product：製品", "primary"), ("Price：価格", "mid"), ("Place：流通", "mid"),
               ("Promotion：販促", "primary")], (2, 2)),
       "4c": ([("Customer Value：顧客価値", "primary"), ("Cost：顧客の負担", "mid"), ("Convenience：利便性", "mid"),
               ("Communication：対話", "primary")], (2, 2)),
       "pest": ([("Politics：政治", "primary"), ("Economy：経済", "mid"), ("Society：社会", "mid"),
                 ("Technology：技術", "primary")], (2, 2)),
       "kpt": ([("Keep：続けること", "primary"), ("Problem：問題", "gray"), ("Try：次に試すこと", "mid")], (1, 3))}


@part("フレームワーク", "定番フレームワーク（kind: 3c / 4p / 4c / pest / kpt）")
def framework_grid(slide, area, kind, contents, size=12):
    """contents: 各ブロックの箇条書き list を順に（3c なら [顧客, 競合, 自社]）"""
    heads, (r, c) = _FW[kind]
    for (t, st), items, a in zip(heads, contents, grid_areas(area, r, c, 12)):
        titled_box(slide, a, t, items, style=st, size=size)


@part("フレームワーク", "5フォース分析（中央の業界と4つの力を矢印で）")
def five_forces(slide, area, center, forces, size=12):
    """forces: {"新規参入の脅威": items, "代替品の脅威": items, "売り手の交渉力": items, "買い手の交渉力": items}"""
    x, y, w, h = area
    bw, bh = w * 0.3, h * 0.3
    cb = (x + (w - bw) / 2, y + (h - bh) / 2, bw, bh)
    pos = {"新規参入の脅威": (x + (w - bw) / 2, y, bw, bh * 0.95), "代替品の脅威": (x + (w - bw) / 2, y + h - bh * 0.95, bw, bh * 0.95),
           "売り手の交渉力": (x, y + (h - bh) / 2, bw * 0.95, bh), "買い手の交渉力": (x + w - bw * 0.95, y + (h - bh) / 2, bw * 0.95, bh)}
    for k, a in pos.items():
        titled_box(slide, a, k, forces.get(k, []), style="gray", size=size, head_h=26)
        connect(slide, a, cb, kind="thick")
    titled_box(slide, cb, "業界内の競争", center, style="primary", size=size, head_h=26)


@part("フレームワーク", "ビジネスモデルキャンバス（9ブロック）")
def business_model_canvas(slide, area, blocks, size=10):
    """blocks: {"KP": [], "KA": [], "KR": [], "VP": [], "CR": [], "CH": [], "CS": [], "C$": [], "R$": []}"""
    names = {"KP": "パートナー", "KA": "主要活動", "KR": "リソース", "VP": "価値提案", "CR": "顧客との関係",
             "CH": "チャネル", "CS": "顧客セグメント", "C$": "コスト構造", "R$": "収益の流れ"}
    x, y, w, h = area
    top, bot = split_rows(area, [2.6, 1], 4)
    c = split_cols(top, [1] * 5, 4)
    half = lambda a: split_rows(a, [1, 1], 4)  # noqa: E731
    lay = {"KP": c[0], "KA": half(c[1])[0], "KR": half(c[1])[1], "VP": c[2], "CR": half(c[3])[0],
           "CH": half(c[3])[1], "CS": c[4]}
    b = split_cols(bot, [1, 1], 4)
    lay["C$"], lay["R$"] = b[0], b[1]
    for k, (bx, by, bw, bh) in lay.items():
        _rect(slide, bx, by, bw, bh, fill=COLOR_PRIMARY_PALE if k == "VP" else COLOR_WHITE, line=COLOR_PRIMARY_LIGHT)
        _textbox(slide, bx + 4, by + 2, bw - 8, 20, names[k], size=11, bold=True, color=COLOR_PRIMARY, dense=True,
                 where=f"bmc見出し[{k}]")
        if blocks.get(k):
            _textbox(slide, bx + 4, by + 22, bw - 8, bh - 26, blocks[k], size=size, bullets=True, dense=True,
                     where=f"bmc[{k}]")


# ---------------------------------------------------------------- 図解・スライド枠
@part("図解", "パイプライン（アイコンのノードを矢印とラベルでつなぐ。データの流れなど）")
def pipeline(slide, area, nodes, edge_labels=None, size=12):
    """nodes: [(icon_kind, ラベル)]  edge_labels: ノード間の矢印に添える文字（len = ノード数-1）"""
    x, y, w, h = area
    cols = split_cols(area, [1] * len(nodes), 50)
    for i, ((k, lab), (cx, cy, cw, ch)) in enumerate(zip(nodes, cols)):
        nw = min(cw, 120)
        icon_label(slide, (cx + (cw - nw) / 2, cy, nw, ch), k, lab, size=size)
        if i < len(nodes) - 1:
            s = min(nw * 0.7, ch - 26)
            ay = cy + s / 2
            _arrow_line(slide, cx + cw + 4, ay, cx + cw + 46, ay, COLOR_PRIMARY_MID, 2)
            if edge_labels and edge_labels[i]:
                _textbox(slide, cx + cw - 10, ay - 26, 70, 22, edge_labels[i], size=10, color=COLOR_PRIMARY_MID,
                         align="center", dense=True, where="pipelineラベル")


@part("スライド枠", "表紙（紺の帯に白抜きタイトル）")
def cover_band(prs, title, subtitle="", date_text="", org="", recipient=""):
    s = _blank(prs)
    _STATE["cover"] = True
    if recipient:
        _textbox(s, 50, 40, 600, 24, recipient, size=14, name="chrome:recipient")
    _rect(s, 0, 150, SLIDE_W, 170, fill=COLOR_PRIMARY, name="chrome:band")
    _textbox(s, 50, 165, 860, 100, title, size=28, bold=True, color=COLOR_WHITE, anchor="middle",
             name="chrome:cover_title", where="表紙タイトル")
    if subtitle:
        _textbox(s, 50, 270, 860, 36, subtitle, size=18, color=COLOR_PRIMARY_PALE, name="chrome:subtitle",
                 where="表紙サブタイトル")
    meta = [t for t in (date_text, org) if t]
    if meta:
        _textbox(s, 50, 400, 600, 60, meta, size=16, name="chrome:meta")
    return s


@part("スライド枠", "章の区切り（左に紺のパネル、右に節の一覧）")
def section_split(prs, title, items=None):
    s = _blank(prs)
    _rect(s, 0, 0, 360, SLIDE_H, fill=COLOR_PRIMARY, name="chrome:panel")
    _textbox(s, 36, 180, 300, 160, title, size=26, bold=True, color=COLOR_WHITE, anchor="middle",
             name="chrome:section_title", where="セクションタイトル")
    if items:
        for i, (rx, ry, rw, rh) in enumerate(split_rows((410, 150, 500, min(48 * len(items), 300)), [1] * len(items), 4)):
            _textbox(s, rx, ry, rw, rh, items[i], size=18, color=COLOR_PRIMARY, anchor="middle", where="セクション項目")
            _line(s, rx, ry + rh + 2, rx + rw, ry + rh + 2, COLOR_GRAY_LIGHT, LINE_W)
    _chrome_footer(prs, s)
    return s
