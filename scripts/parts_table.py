# -*- coding: utf-8 -*-
"""parts_table - 表の部品（評価表・マトリクス・RACI・ヒートマップなど）"""
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

from pptx_helpers import *  # noqa: F401,F403
from pptx_helpers import (check_tokens, _ALIGN, _add_runs, _box, _bullet, _cell_border, _check_size, _circle, _font,
                          _line, _norm_items, _rect, _rgb, _split, _textbox, estimate_height, part)
from parts_text import split_cols

_MARKS = {"◎": COLOR_PRIMARY, "○": COLOR_PRIMARY_MID, "△": COLOR_GRAY_DARK, "×": COLOR_ACCENT,
          "✓": COLOR_PRIMARY_MID}


def _native_table(slide, area, data, col_ratios, style, size=12, stretch=False, where="table"):
    """汎用ネイティブ表。style(i, j, cell) -> (fill, color, bold, align)"""
    x, y, w, h = area
    _check_size(size, True, where)
    nr, nc = len(data), len(data[0])
    widths = _split(w, col_ratios or [1] * nc, 0)
    mx, my = 5, 3
    heights = []
    for r in data:
        for j, c in enumerate(r):
            for lv, tt in _norm_items(c if c is not None else ""):
                check_tokens(tt, widths[j] - mx * 2, size, where)
        need = max(estimate_height(c if c is not None else "", widths[j] - mx * 2, size + 1,
                                   isinstance(c, list)) for j, c in enumerate(r))
        heights.append(max(need + my * 2 + 2, size * 2.0))
    total = sum(heights)
    if total > h + 0.5:
        raise ValueError(f"[{where}] 表が領域に収まりません: 推定 {total:.0f}pt > {h:.0f}pt。行・文字を減らすか表を分割してください。")
    if stretch:
        heights = [hh + (h - total) / nr for hh in heights]
    gf = slide.shapes.add_table(nr, nc, Pt(x), Pt(y), Pt(w), Pt(sum(heights)))
    tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Pt(cw)
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Pt(hh)
    for i, r in enumerate(data):
        for j, c in enumerate(r):
            fill, col, bold, al = style(i, j, c)
            cell = tbl.cell(i, j)
            _cell_border(cell, COLOR_GRAY_DARK, 0.5)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            cell.margin_left = cell.margin_right = Pt(mx)
            cell.margin_top = cell.margin_bottom = Pt(my)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            for k, (lv, t) in enumerate(_norm_items(c if c is not None else "")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = _ALIGN[al]
                if isinstance(c, list):
                    _bullet(p, lv, size)
                if i > 0 and t[:1] in _MARKS and (len(t) == 1 or t[1] == " "):   # "◎ 説明" 形式のみ
                    r0 = p.add_run()
                    r0.text = t[0]
                    _font(r0.font, size + 2, _MARKS[t[0]], True)
                    t = t[1:]
                if t:
                    _add_runs(p, t, size, col, bold)
    return gf


def _hdr(i, j, c):
    return COLOR_PRIMARY, COLOR_WHITE, True, "center"


@part("表", "比較評価表（◎○△×を色付きで表示、推奨列を強調）")
def comparison_table(slide, area, criteria, options, cells, recommend=None, col_ratios=None, size=12, stretch=False):
    """cells[行=評価項目][列=選択肢]。セル先頭の ◎○△× は自動で色付け（例: "◎ 低コスト"）"""
    data = [["評価項目"] + list(options)] + [[c] + list(r) for c, r in zip(criteria, cells)]

    def st(i, j, c):
        rec = recommend is not None and j == recommend + 1
        if i == 0:
            return (COLOR_PRIMARY_MID if rec else COLOR_PRIMARY), COLOR_WHITE, True, "center"
        if j == 0:
            return COLOR_GRAY_PALE, COLOR_TEXT, True, "left"
        return (COLOR_PRIMARY_PALE if rec else COLOR_WHITE), COLOR_TEXT, False, "left"
    return _native_table(slide, area, data, col_ratios or [1.2] + [1] * len(options), st, size,
                         stretch=stretch, where="comparison_table")


@part("表", "マトリクス表（行見出し×列見出し、カード型）")
def matrix_table(slide, area, row_labels, col_labels, cells, corner="", label_ratio=0.8, size=13):
    x, y, w, h = area
    gap = 4
    cols = split_cols(area, [label_ratio] + [1] * len(col_labels), gap)
    hh = 34
    rows_h = _split(h - hh - gap, [1] * len(row_labels), gap)
    if corner:
        _box(slide, cols[0][0], y, cols[0][2], hh, corner, style="gray", size=size)
    for j, lab in enumerate(col_labels):
        _box(slide, cols[j + 1][0], y, cols[j + 1][2], hh, lab, style="label2", size=size,
             where=f"matrix_table列見出し[{j}]")
    cy = y + hh + gap
    for i, lab in enumerate(row_labels):
        _box(slide, cols[0][0], cy, cols[0][2], rows_h[i], lab, style="label", size=size,
             where=f"matrix_table行見出し[{i}]")
        for j in range(len(col_labels)):
            _box(slide, cols[j + 1][0], cy, cols[j + 1][2], rows_h[i], cells[i][j], style="text", size=size,
                 where=f"matrix_table[{i},{j}]")
        cy += rows_h[i] + gap


@part("表", "スコア表（●○の点数表示＋合計、重み付け可）")
def score_table(slide, area, criteria, options, scores, max_score=5, weights=None, size=12, stretch=False):
    weights = weights or [1] * len(criteria)
    head = ["評価項目"] + (["重み"] if any(wt != 1 for wt in weights) else []) + list(options)
    data = [head]
    for c, wt, row in zip(criteria, weights, scores):
        data.append([c] + ([f"×{wt}"] if len(head) > len(options) + 1 else [])
                    + ["●" * v + "○" * (max_score - v) for v in row])
    totals = [sum(wt * row[k] for wt, row in zip(weights, scores)) for k in range(len(options))]
    best = max(totals)
    data.append(["合計"] + ([""] if len(head) > len(options) + 1 else []) + [str(t) for t in totals])
    off = len(head) - len(options)

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if i == len(data) - 1:
            win = j >= off and totals[j - off] == best
            return (COLOR_PRIMARY_PALE if win else COLOR_GRAY_PALE), COLOR_PRIMARY, True, "center"
        if j == 0:
            return COLOR_GRAY_PALE, COLOR_TEXT, True, "left"
        return COLOR_WHITE, COLOR_PRIMARY_MID, False, "center"
    return _native_table(slide, area, data, [1.6] + [0.5] * (off - 1) + [1] * len(options), st, size,
                         stretch=stretch, where="score_table")


def _harvey(slide, cx, cy, d, v):
    """ハーベイボール（v=0..4）"""
    _rect(slide, cx - d / 2, cy - d / 2, d, d, fill=COLOR_WHITE, line=COLOR_PRIMARY, line_w=1.25,
          shape=MSO_SHAPE.OVAL)
    if v >= 4:
        _rect(slide, cx - d / 2, cy - d / 2, d, d, fill=COLOR_PRIMARY, shape=MSO_SHAPE.OVAL)
    elif v > 0:
        sp = _rect(slide, cx - d / 2, cy - d / 2, d, d, fill=COLOR_PRIMARY, shape=MSO_SHAPE.PIE)
        sp.adjustments[0] = 270.0 * 60000 / 100000
        sp.adjustments[1] = ((270 + 90 * v) % 360) * 60000 / 100000


@part("表", "ハーベイボール表（充足度を ○◔◑◕● で表示）")
def harvey_table(slide, area, row_labels, col_labels, values, label_ratio=1.6, size=12):
    """values[行][列] = 0..4"""
    x, y, w, h = area
    cols = split_cols(area, [label_ratio] + [1] * len(col_labels), 0)
    hh = 30
    for j, lab in enumerate(col_labels):
        _box(slide, cols[j + 1][0] + 2, y, cols[j + 1][2] - 4, hh, lab, style="label", size=size,
             dense=True, where=f"harvey_table列[{j}]")
    rh = (h - hh - 4) / len(row_labels)
    d = min(22, rh - 8)
    for i, lab in enumerate(row_labels):
        ry = y + hh + 4 + i * rh
        _textbox(slide, x, ry, cols[0][2] - 6, rh, lab, size=size, bold=True, anchor="middle", dense=True,
                 where=f"harvey_table行[{i}]")
        for j in range(len(col_labels)):
            _harvey(slide, cols[j + 1][0] + cols[j + 1][2] / 2, ry + rh / 2, d, values[i][j])
        _line(slide, x, ry + rh, x + w, ry + rh, COLOR_GRAY_LIGHT, LINE_W)


_RACI = {"R": (COLOR_PRIMARY, COLOR_WHITE), "A": (COLOR_PRIMARY_MID, COLOR_WHITE),
         "C": (COLOR_PRIMARY_PALE, COLOR_PRIMARY), "I": (COLOR_GRAY_PALE, COLOR_TEXT_SUB)}


@part("表", "RACI 表（役割分担 R/A/C/I を色分け）")
def raci_table(slide, area, tasks, roles, matrix, size=12, stretch=False):
    data = [["タスク"] + list(roles)] + [[t] + list(r) for t, r in zip(tasks, matrix)]

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if j == 0:
            return COLOR_WHITE, COLOR_TEXT, False, "left"
        f, col = _RACI.get((c or "").strip()[:1], (COLOR_WHITE, COLOR_TEXT))
        return f, col, True, "center"
    return _native_table(slide, area, data, [2.4] + [1] * len(roles), st, size, stretch=stretch, where="raci_table")


@part("表", "項目／内容の2列表（概要表・諸元表）")
def key_value_table(slide, area, pairs, key_ratio=0.25, size=13, stretch=False):
    data = [[k, v] for k, v in pairs]

    def st(i, j, c):
        return (COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, "left") if j == 0 else (COLOR_WHITE, COLOR_TEXT, False, "left")
    return _native_table(slide, area, data, [key_ratio, 1 - key_ratio], st, size, stretch=stretch,
                         where="key_value_table")


@part("表", "ヒートマップ表（値の大小を5段階の青で塗る）")
def heatmap_table(slide, area, row_labels, col_labels, values, fmt="{:,.0f}", size=12, stretch=False):
    flat = [v for r in values for v in r]
    lo, hi = min(flat), max(flat)
    ramp = [(COLOR_WHITE, COLOR_TEXT), (COLOR_PRIMARY_PALE, COLOR_TEXT), (COLOR_PRIMARY_LIGHT, COLOR_TEXT),
            (COLOR_PRIMARY_MID, COLOR_WHITE), (COLOR_PRIMARY, COLOR_WHITE)]
    data = [[""] + list(col_labels)] + [[lab] + [fmt.format(v) for v in r] for lab, r in zip(row_labels, values)]

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if j == 0:
            return COLOR_GRAY_PALE, COLOR_TEXT, True, "left"
        v = values[i - 1][j - 1]
        k = 0 if hi == lo else min(4, int((v - lo) / (hi - lo) * 5))
        f, col = ramp[k]
        return f, col, k >= 3, "center"
    return _native_table(slide, area, data, [1.4] + [1] * len(col_labels), st, size, stretch=stretch, where="heatmap_table")


_STATUS = {"完了": (COLOR_PRIMARY, COLOR_WHITE, True), "進行中": (COLOR_PRIMARY_MID, COLOR_WHITE, True),
           "未着手": (COLOR_GRAY_PALE, COLOR_TEXT_SUB, False), "遅延": (COLOR_WHITE, COLOR_ACCENT, True),
           "保留": (COLOR_GRAY_LIGHT, COLOR_TEXT, False)}


@part("表", "進捗・ステータス表（完了/進行中/未着手/遅延/保留を色分け）")
def status_table(slide, area, header, rows, status_col, col_ratios=None, size=12, stretch=False):
    data = [list(header)] + [list(r) for r in rows]

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if j == status_col:
            f, col, b = _STATUS.get(c, (COLOR_WHITE, COLOR_TEXT, False))
            return f, col, b, "center"
        return COLOR_WHITE, COLOR_TEXT, False, "left"
    return _native_table(slide, area, data, col_ratios, st, size, stretch=stretch, where="status_table")


@part("表", "対応表（✓ マーク。機能×対象など）")
def check_table(slide, area, row_labels, col_labels, marks, size=12, stretch=False):
    """marks[行][列] は True/False または文字列"""
    data = [[""] + list(col_labels)] + [[lab] + [("✓" if m is True else "" if m is False else m) for m in r]
                                       for lab, r in zip(row_labels, marks)]

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if j == 0:
            return COLOR_GRAY_PALE, COLOR_TEXT, True, "left"
        return COLOR_WHITE, COLOR_PRIMARY_MID, True, "center"
    return _native_table(slide, area, data, [1.6] + [1] * len(col_labels), st, size, stretch=stretch, where="check_table")


@part("表", "行を強調できる一覧表（強調行は淡青＋太字）")
def highlight_table(slide, area, header, rows, highlight_rows=(), col_ratios=None, align=None, size=12, stretch=False):
    data = [list(header)] + [list(r) for r in rows]
    amap = {"l": "left", "c": "center", "r": "right"}
    align = align or ["l"] * len(header)

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if i - 1 in highlight_rows:
            return COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, amap[align[j]]
        return COLOR_WHITE, COLOR_TEXT, False, amap[align[j]]
    return _native_table(slide, area, data, col_ratios, st, size, stretch=stretch, where="highlight_table")


@part("表", "宿題・アクションアイテム一覧（内容・担当・期限・状況）")
def action_items(slide, area, items, size=12, stretch=False):
    """items: [(内容, 担当, 期限, 状況)]  状況は 完了/進行中/未着手/遅延/保留"""
    rows = [[str(i + 1)] + list(it) for i, it in enumerate(items)]
    return status_table(slide, area, ["No", "内容", "担当", "期限", "状況"], rows, status_col=4,
                        col_ratios=[0.6, 4, 1.2, 1.1, 1], size=size, stretch=stretch)
