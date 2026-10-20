# -*- coding: utf-8 -*-
"""parts_connect - 図の基盤: アイコン・矢印の種類・自動配線・番号注釈・凡例"""
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE

from pptx_helpers import *  # noqa: F401,F403
from pptx_helpers import (_arrow_line, _circle, _line, _no_effect, _poly, _put_text, _rect, _rgb, _textbox,
                          part, text_width)
from parts_text import split_rows

# ---------------------------------------------------------------- アイコン（図形の組み合わせ。編集可能なグループ）


def _person(g, x, y, s, col):
    _rect(g, x + 0.31 * s, y + 0.02 * s, 0.38 * s, 0.38 * s, fill=col, shape=MSO_SHAPE.OVAL)
    b = _rect(g, x + 0.14 * s, y + 0.46 * s, 0.72 * s, 0.52 * s, fill=col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    b.adjustments[0] = 0.45


def _draw_icon(g, kind, x, y, s, col):
    W = COLOR_WHITE
    if kind == "person":
        _person(g, x, y, s, col)
    elif kind == "people":
        _person(g, x + 0.34 * s, y, 0.66 * s, COLOR_PRIMARY_LIGHT)
        _person(g, x, y + 0.2 * s, 0.8 * s, col)
    elif kind == "document":
        _rect(g, x + 0.16 * s, y + 0.02 * s, 0.68 * s, 0.96 * s, fill=W, line=col, line_w=max(1.25, s * 0.05),
              shape=MSO_SHAPE.FOLDED_CORNER)
        for i in range(3):
            _line(g, x + 0.28 * s, y + (0.28 + 0.18 * i) * s, x + 0.72 * s, y + (0.28 + 0.18 * i) * s, col,
                  max(1, s * 0.04))
    elif kind == "pc":
        _rect(g, x, y + 0.08 * s, s, 0.62 * s, fill=col)
        _rect(g, x + 0.07 * s, y + 0.14 * s, 0.86 * s, 0.48 * s, fill=W)
        _rect(g, x + 0.42 * s, y + 0.7 * s, 0.16 * s, 0.14 * s, fill=col)
        _rect(g, x + 0.24 * s, y + 0.84 * s, 0.52 * s, 0.08 * s, fill=col)
    elif kind == "server":
        for i in range(3):
            _rect(g, x + 0.1 * s, y + i * 0.34 * s, 0.8 * s, 0.28 * s, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            _rect(g, x + 0.7 * s, y + i * 0.34 * s + 0.1 * s, 0.08 * s, 0.08 * s, fill=W, shape=MSO_SHAPE.OVAL)
    elif kind == "database":
        _rect(g, x + 0.14 * s, y, 0.72 * s, s, fill=col, shape=MSO_SHAPE.CAN)
    elif kind == "cloud":
        _rect(g, x, y + 0.16 * s, s, 0.68 * s, fill=col, shape=MSO_SHAPE.CLOUD)
    elif kind == "building":
        _rect(g, x + 0.16 * s, y + 0.04 * s, 0.68 * s, 0.94 * s, fill=col)
        for r in range(3):
            for c in range(3):
                _rect(g, x + (0.25 + 0.18 * c) * s, y + (0.14 + 0.2 * r) * s, 0.12 * s, 0.12 * s, fill=W)
        _rect(g, x + 0.42 * s, y + 0.76 * s, 0.16 * s, 0.22 * s, fill=W)
    elif kind == "factory":
        _poly(g, [(x, y + s), (x, y + 0.45 * s), (x + 0.3 * s, y + 0.62 * s), (x + 0.3 * s, y + 0.45 * s),
                  (x + 0.6 * s, y + 0.62 * s), (x + 0.6 * s, y + 0.1 * s), (x + 0.78 * s, y + 0.1 * s),
                  (x + 0.78 * s, y + 0.45 * s), (x + s, y + 0.45 * s), (x + s, y + s)], fill=col)
        for c in range(3):
            _rect(g, x + (0.1 + 0.3 * c) * s, y + 0.74 * s, 0.14 * s, 0.12 * s, fill=W)
    elif kind == "money":
        _circle(g, x + 0.04 * s, y + 0.04 * s, 0.92 * s, "¥", fill=col, size=s * 0.48)
    elif kind == "phone":
        _rect(g, x + 0.26 * s, y, 0.48 * s, s, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(g, x + 0.31 * s, y + 0.1 * s, 0.38 * s, 0.7 * s, fill=W)
        _rect(g, x + 0.45 * s, y + 0.85 * s, 0.1 * s, 0.1 * s, fill=W, shape=MSO_SHAPE.OVAL)
    elif kind == "mail":
        _rect(g, x, y + 0.18 * s, s, 0.64 * s, fill=W, line=col, line_w=max(1.25, s * 0.05))
        _line(g, x, y + 0.18 * s, x + 0.5 * s, y + 0.55 * s, col, max(1.25, s * 0.05))
        _line(g, x + s, y + 0.18 * s, x + 0.5 * s, y + 0.55 * s, col, max(1.25, s * 0.05))
    elif kind == "gear":
        _rect(g, x + 0.04 * s, y + 0.04 * s, 0.92 * s, 0.92 * s, fill=col, shape=MSO_SHAPE.GEAR_6)
    elif kind == "lock":
        _rect(g, x + 0.27 * s, y + 0.04 * s, 0.46 * s, 0.62 * s, line=col, line_w=max(2, s * 0.09),
              shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
        _rect(g, x + 0.14 * s, y + 0.42 * s, 0.72 * s, 0.56 * s, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(g, x + 0.45 * s, y + 0.58 * s, 0.1 * s, 0.22 * s, fill=W)
    elif kind == "chart":
        _line(g, x + 0.06 * s, y + 0.06 * s, x + 0.06 * s, y + 0.94 * s, col, max(1.25, s * 0.05))
        _line(g, x + 0.06 * s, y + 0.94 * s, x + 0.96 * s, y + 0.94 * s, col, max(1.25, s * 0.05))
        for i, hh in enumerate((0.35, 0.55, 0.8)):
            _rect(g, x + (0.18 + 0.27 * i) * s, y + (0.9 - hh) * s, 0.18 * s, hh * s, fill=col)
    elif kind == "check":
        _circle(g, x + 0.04 * s, y + 0.04 * s, 0.92 * s, "✓", fill=col, size=s * 0.5)
    elif kind == "warning":
        _rect(g, x, y + 0.04 * s, s, 0.9 * s, fill=COLOR_STATUS_WARN, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        _textbox(g, x + 0.3 * s, y + 0.3 * s, 0.4 * s, 0.6 * s, "!", size=max(8, s * 0.45), bold=True,
                 color=COLOR_WHITE, align="center", anchor="middle", check=False, exempt=True, name="deco:icon")
    elif kind == "clock":
        _rect(g, x + 0.04 * s, y + 0.04 * s, 0.92 * s, 0.92 * s, fill=W, line=col, line_w=max(1.5, s * 0.07),
              shape=MSO_SHAPE.OVAL)
        _line(g, x + 0.5 * s, y + 0.5 * s, x + 0.5 * s, y + 0.2 * s, col, max(1.5, s * 0.06))
        _line(g, x + 0.5 * s, y + 0.5 * s, x + 0.72 * s, y + 0.6 * s, col, max(1.5, s * 0.06))
    else:
        raise ValueError(f"[icon] 未知のアイコン「{kind}」。使えるもの: {', '.join(ICONS)}")


ICONS = ["person", "people", "document", "pc", "server", "database", "cloud", "building", "factory", "money",
         "phone", "mail", "gear", "lock", "chart", "check", "warning", "clock"]


@part("アイコン", "アイコン（図形で描画・編集可能）。kind は ICONS の18種")
def icon(slide, x, y, s, kind, color=COLOR_PRIMARY):
    """kind: person / people / document / pc / server / database / cloud / building / factory / money /
    phone / mail / gear / lock / chart / check / warning / clock"""
    g = slide.shapes.add_group_shape()
    _draw_icon(g, kind, x, y, s, color)
    g.name = f"icon:{kind}"
    return g


@part("アイコン", "アイコン＋ラベル（＋補足）を縦に並べたノード。外枠の領域を返す")
def icon_label(slide, area, kind, label, sub=None, color=COLOR_PRIMARY, size=12):
    x, y, w, h = area
    th = 22 + (18 if sub else 0)
    s = min(w * 0.7, h - th - 4)
    icon(slide, x + (w - s) / 2, y, s, kind, color)
    _textbox(slide, x, y + s + 3, w, 22, label, size=size, bold=True, align="center", dense=True,
             where="icon_label")
    if sub:
        _textbox(slide, x, y + s + 23, w, 18, sub, size=10, color=COLOR_TEXT_SUB, align="center",
                 dense=True, where="icon_label補足")
    return area


@part("アイコン", "アイコン付きのステップ（アイコン → 矢印 → アイコン）")
def icon_steps(slide, area, steps, size=12):
    """steps: [(icon_kind, 見出し, 説明)]"""
    from parts_text import split_cols
    x, y, w, h = area
    cols = split_cols(area, [1] * len(steps), 36)
    s = min(cols[0][2] * 0.5, h * 0.42, 70)
    for i, ((k, t, d), (cx, cy, cw, ch)) in enumerate(zip(steps, cols)):
        _circle(slide, cx + (cw - s * 1.3) / 2, cy, s * 1.3, None, fill=COLOR_PRIMARY_PALE)
        icon(slide, cx + (cw - s * 0.8) / 2, cy + s * 0.25, s * 0.8, k)
        _textbox(slide, cx, cy + s * 1.3 + 6, cw, 24, f"{i + 1}. {t}", size=size + 1, bold=True,
                 color=COLOR_PRIMARY, align="center", where=f"icon_steps見出し[{i}]")
        if d:
            _textbox(slide, cx, cy + s * 1.3 + 32, cw, ch - s * 1.3 - 32, d, size=size, align="center",
                     where=f"icon_steps説明[{i}]")
        if i < len(steps) - 1:
            _arrow_line(slide, cx + cw + 6, cy + s * 0.65, cx + cw + 30, cy + s * 0.65, COLOR_PRIMARY_MID, 2.5)


# ---------------------------------------------------------------- 線の種類・配線
FLOWS = {  # 種類: (色, 太さ, 破線, 両矢印, 凡例名)
    "solid": (COLOR_PRIMARY, 1.5, None, False, "流れ"),
    "dashed": (COLOR_PRIMARY_MID, 1.5, MSO_LINE_DASH_STYLE.DASH, False, "間接的な流れ"),
    "dotted": (COLOR_GRAY_DARK, 1.25, MSO_LINE_DASH_STYLE.ROUND_DOT, False, "参照・情報"),
    "thick": (COLOR_PRIMARY, 3.0, None, False, "主要な流れ"),
    "both": (COLOR_PRIMARY, 1.5, None, True, "相互"),
    "service": (COLOR_PRIMARY, 1.75, None, False, "サービス"),
    "money": (COLOR_PRIMARY_MID, 1.75, MSO_LINE_DASH_STYLE.DASH, False, "お金"),
    "goods": (COLOR_GRAY_DARK, 3.0, None, False, "モノ"),
    "info": (COLOR_GRAY_DARK, 1.25, MSO_LINE_DASH_STYLE.ROUND_DOT, False, "情報"),
}


def _xywh(a):
    if isinstance(a, (tuple, list)):
        return tuple(a)
    return (a.left / 12700, a.top / 12700, a.width / 12700, a.height / 12700)


def _styled(slide, x1, y1, x2, y2, kind, head=True):
    col, wd, dash, both, _ = FLOWS[kind]
    ln = _arrow_line(slide, x1, y1, x2, y2, col, wd, head=head, both=both and head)
    if dash is not None:
        ln.line.dash_style = dash
    return ln


def _edge_label(slide, x, y, label, kind):
    lw = min(150, text_width(label, 10) + 14)
    tb = _textbox(slide, x - lw / 2, y - 10, lw, 20, label, size=10, color=FLOWS[kind][0], align="center",
                  anchor="middle", dense=True, where="線ラベル")
    tb.fill.solid()
    tb.fill.fore_color.rgb = _rgb(COLOR_WHITE)
    tb.name = "deco:edge_label"
    return tb


def _edge(a, tx, ty):
    x, y, w, h = a
    cx, cy = x + w / 2, y + h / 2
    dx, dy = tx - cx, ty - cy
    if dx == 0 and dy == 0:
        return cx, cy
    k = min((w / 2) / abs(dx) if dx else 1e9, (h / 2) / abs(dy) if dy else 1e9)
    return cx + dx * k, cy + dy * k


@part("線・矢印", "2つの箱を縁から縁へ自動でつなぐ（種類: solid/dashed/dotted/thick/both/service/money/goods/info）")
def connect(slide, a, b, kind="solid", label=None, offset=0):
    """a, b は (x,y,w,h) か図形。offset で平行線をずらす（往復の矢印を2本引くとき）"""
    a, b = _xywh(a), _xywh(b)
    bc = (b[0] + b[2] / 2, b[1] + b[3] / 2)
    ac = (a[0] + a[2] / 2, a[1] + a[3] / 2)
    sx, sy = _edge(a, *bc)
    ex, ey = _edge(b, *ac)
    if offset:
        import math
        L = math.hypot(ex - sx, ey - sy) or 1
        nx, ny = -(ey - sy) / L * offset, (ex - sx) / L * offset
        sx, sy, ex, ey = sx + nx, sy + ny, ex + nx, ey + ny
    ln = _styled(slide, sx, sy, ex, ey, kind)
    if label:
        _edge_label(slide, (sx + ex) / 2, (sy + ey) / 2, label, kind)
    return ln


@part("線・矢印", "カギ線で2つの箱をつなぐ（first='h' 横→縦→横 / 'v' 縦→横→縦）")
def elbow_connect(slide, a, b, first="h", kind="solid", label=None):
    a, b = _xywh(a), _xywh(b)
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    if first == "h":
        right = bx > ax + aw / 2
        sx, sy = (ax + aw, ay + ah / 2) if right else (ax, ay + ah / 2)
        ex, ey = (bx, by + bh / 2) if right else (bx + bw, by + bh / 2)
        mx = (sx + ex) / 2
        pts = [(sx, sy), (mx, sy), (mx, ey), (ex, ey)]
    else:
        down = by > ay + ah / 2
        sx, sy = (ax + aw / 2, ay + ah) if down else (ax + aw / 2, ay)
        ex, ey = (bx + bw / 2, by) if down else (bx + bw / 2, by + bh)
        my = (sy + ey) / 2
        pts = [(sx, sy), (sx, my), (ex, my), (ex, ey)]
    for i in range(3):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        if abs(x1 - x2) > 0.1 or abs(y1 - y2) > 0.1:
            _styled(slide, x1, y1, x2, y2, kind, head=i == 2)
    if label:
        (x1, y1), (x2, y2) = pts[1], pts[2]
        _edge_label(slide, (x1 + x2) / 2, (y1 + y2) / 2, label, kind)


@part("線・矢印", "曲線の矢印（2点間）")
def curved_arrow(slide, p1, p2, kind="solid", label=None):
    col, wd, dash, both, _ = FLOWS[kind]
    ln = slide.shapes.add_connector(MSO_CONNECTOR.CURVE, Pt(p1[0]), Pt(p1[1]), Pt(p2[0]), Pt(p2[1]))
    ln.line.color.rgb = _rgb(col)
    ln.line.width = Pt(wd)
    if dash is not None:
        ln.line.dash_style = dash
    lnx = ln.line._get_or_add_ln()
    lnx.append(lnx.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    _no_effect(ln)
    ln.name = "arrow"
    if label:
        _edge_label(slide, (p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2, label, kind)
    return ln


@part("線・矢印", "循環の矢印（ぐるっと回る矢印＋中央ラベル）")
def loop_arrow(slide, area, label=None, size=13):
    x, y, w, h = area
    d = min(w, h)
    sp = _rect(slide, x + (w - d) / 2, y + (h - d) / 2, d, d, fill=COLOR_PRIMARY_LIGHT, shape=MSO_SHAPE.CIRCULAR_ARROW)
    sp.name = "arrow:loop"
    if label:
        _textbox(slide, x + (w - d * 0.6) / 2, y + (h - d * 0.4) / 2, d * 0.6, d * 0.4, label, size=size, bold=True,
                 color=COLOR_PRIMARY, align="center", anchor="middle", where="loop_arrow")
    return sp


_BLOCK = {"right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW, "up": MSO_SHAPE.UP_ARROW,
          "down": MSO_SHAPE.DOWN_ARROW, "both": MSO_SHAPE.LEFT_RIGHT_ARROW}


@part("線・矢印", "文字入りの太い矢印（right/left/up/down/both）")
def block_arrow(slide, area, body, direction="right", style="pale", size=14):
    x, y, w, h = area
    fill, col = {"pale": (COLOR_PRIMARY_PALE, COLOR_PRIMARY), "primary": (COLOR_PRIMARY, COLOR_WHITE),
                 "mid": (COLOR_PRIMARY_MID, COLOR_WHITE), "gray": (COLOR_GRAY_LIGHT, COLOR_TEXT)}[style]
    sp = _rect(slide, x, y, w, h, fill=fill, shape=_BLOCK[direction])
    sp.name = "arrow:block"
    if direction in ("right", "left", "both"):
        head = min(h * 0.5, w * 0.3)
        tx = x + (head if direction in ("left", "both") else 4)
        tw = w - head * (2 if direction == "both" else 1) - 4
        box = (tx, y + h * 0.25, tw, h * 0.5)
    else:
        head = min(w * 0.5, h * 0.3)
        box = (x + w * 0.25, y + (head if direction == "up" else 4), w * 0.5, h - head - 4)
    _textbox(slide, *box, body, size=size, bold=True, color=col, align="center", anchor="middle",
             where="block_arrow")
    return sp


@part("線・矢印", "線の種類の凡例（connect の kind を並べる）。幅を返す")
def flow_legend(slide, x, y, kinds, size=10):
    cx = x
    for k in kinds:
        _styled(slide, cx, y + 9, cx + 30, y + 9, k)
        lab = FLOWS[k][4]
        tw = text_width(lab, size) + 8
        _textbox(slide, cx + 34, y, tw, 18, lab, size=size, dense=True, anchor="middle", where="flow_legend",
                 check=False)
        cx += 34 + tw + 12
    return cx - x


@part("線・矢印", "図の上に番号を打ち、欄外に説明を並べる（任意の図に使える番号注釈）")
def annotate_points(slide, points, notes, area, size=12, style="alert"):
    """points: [(x, y)] 絶対座標(pt)。notes: 説明（points と同数）。area: 説明を並べる領域"""
    fill = {"alert": COLOR_ACCENT, "primary": COLOR_PRIMARY, "mid": COLOR_PRIMARY_MID}[style]
    for i, (px, py) in enumerate(points):
        _circle(slide, px - 11, py - 11, 22, i + 1, fill=fill, size=11)
    for i, (n, (x, y, w, h)) in enumerate(zip(notes, split_rows(area, [1] * len(notes), 6))):
        _circle(slide, x, y + 2, 22, i + 1, fill=fill, size=11)
        _textbox(slide, x + 30, y, w - 30, h, n, size=size, where=f"annotate_points[{i}]")


@part("線・矢印", "範囲を示すかっこ線（期間・対象範囲にラベル）")
def span_bracket(slide, x1, x2, y, label, below=True, size=11):
    t = 8 if below else -8
    _line(slide, x1, y, x1, y + t, COLOR_PRIMARY, 1.25)
    _line(slide, x2, y, x2, y + t, COLOR_PRIMARY, 1.25)
    _line(slide, x1, y + t, x2, y + t, COLOR_PRIMARY, 1.25)
    ly = y + t + 2 if below else y + t - 22
    _textbox(slide, x1, ly, x2 - x1, 20, label, size=size, bold=True, color=COLOR_PRIMARY, align="center",
             dense=True, where="span_bracket")
