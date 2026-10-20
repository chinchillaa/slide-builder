# -*- coding: utf-8 -*-
"""parts_diagram - 報告書で多い図解の型（ゾーン図・ビジネスモデル図・仕様書型本文・論点マップなど）"""
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from pptx_helpers import *  # noqa: F401,F403
from pptx_helpers import (_add_runs, _arrow_line, _box, _circle, _font, _line, _norm_items, _put_text, _rect,
                          _rgb, _set_ea, _split, _textbox, _triangle, check_tokens, estimate_height,
                          estimate_lines, part, text_width)
from parts_connect import FLOWS, connect, flow_legend, icon
from parts_structure import logic_tree
from parts_table import _native_table
from parts_text import split_cols, split_rows, tag, titled_box


@part("図解", "ゾーン（点線枠で領域をまとめ、左上にラベル）。内側の領域を返す")
def zone(slide, area, label, style="dashed", size=12):
    """style: dashed（点線・塗りなし） / pale（淡青の塗り） / gray（薄グレーの塗り） / solid（紺の実線）"""
    x, y, w, h = area
    fill, line, dash = {"dashed": (None, COLOR_GRAY_DARK, True), "pale": (COLOR_PRIMARY_PALE, None, False),
                        "gray": (COLOR_GRAY_PALE, COLOR_GRAY, False), "solid": (None, COLOR_PRIMARY, False)}[style]
    sp = _rect(slide, x, y, w, h, fill=fill, line=line, line_w=1.25 if style == "solid" else LINE_W)
    if dash:
        sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sp.name = "zone"
    tag(slide, x + 8, y + 6, label, style="primary" if style != "pale" else "outline", size=size - 1)
    return (x + 8, y + 32, w - 16, h - 40)


@part("図解", "ビジネスモデル図（関係者をアイコンで置き、お金・サービス・モノ・情報の流れでつなぐ）")
def business_model(slide, area, actors, flows, columns=None, legend=True, size=11):
    """actors: {id: (ラベル, 列, 行, icon_kind)}  列・行は 0 始まりの格子位置
    flows: [(from_id, to_id, 種類, ラベル)] 種類は service / money / goods / info / solid / dashed …
    columns: 列見出し（例: ["ユーザー", "事業者", "メーカー"]）"""
    x, y, w, h = area
    n_cols = max(a[1] for a in actors.values()) + 1
    n_rows = max(a[2] for a in actors.values()) + 1
    top = y
    if columns:
        for j, (cx, cy, cw, ch) in enumerate(split_cols((x, y, w, 24), [1] * n_cols, 8)):
            _box(slide, cx, cy, cw, ch, columns[j], style="pale", size=size, dense=True, where=f"business_model列[{j}]")
        top = y + 32
    bottom = y + h - (24 if legend else 0)
    cw = w / n_cols
    rh = (bottom - top) / n_rows
    nw, nh = min(cw * 0.6, 130), min(rh * 0.75, 96)
    boxes = {}
    for aid, (lab, c, r, kind) in actors.items():
        bx = x + c * cw + (cw - nw) / 2
        by = top + r * rh + (rh - nh) / 2
        s = min(nh - 26, nw * 0.5)
        if kind:
            icon(slide, bx + (nw - s) / 2, by, s, kind)
        _box(slide, bx, by + s + 4, nw, nh - s - 4, lab, style="text", size=size, align="center", dense=True,
             where=f"business_model[{aid}]", inset=(2, 1))
        boxes[aid] = (bx, by, nw, nh)
    total = {}
    for f in flows:
        key = tuple(sorted(f[:2]))
        total[key] = total.get(key, 0) + 1
    seen = {}
    for a, b, kind, lab in flows:
        key = tuple(sorted((a, b)))
        k = seen.get(key, 0)
        seen[key] = k + 1
        # 同じ2者間に複数の流れがあるときは平行にずらす（向きに関係なく同じ側へ揃える）
        off = (k - (total[key] - 1) / 2) * 14 * (1 if (a, b) == key else -1)
        connect(slide, boxes[a], boxes[b], kind=kind, label=lab, offset=off)
    if legend:
        kinds = list(dict.fromkeys(f[2] for f in flows))
        flow_legend(slide, x + w - 110 * len(kinds), y + h - 20, kinds)


@part("図解", "業務フローの Before / After（縦の手順を左右に並べ、変わった手順を強調）")
def flow_before_after(slide, area, before, after, changed=(), labels=("これまで", "これから"), size=12):
    """before/after: 手順の list。changed: after 側で変わった手順の index"""
    x, y, w, h = area
    a, b = split_cols(area, [1, 1], 60)
    for col, steps, lab, is_after in ((a, before, labels[0], False), (b, after, labels[1], True)):
        cx, cy, cw, ch = col
        _box(slide, cx, cy, cw, 28, lab, style="label" if is_after else "gray", size=size + 1)
        rows = split_rows((cx + cw * 0.1, cy + 38, cw * 0.8, ch - 38), [1] * len(steps), 18)
        for i, (st, (rx, ry, rw, rh)) in enumerate(zip(steps, rows)):
            on = is_after and i in changed
            _box(slide, rx, ry, rw, rh, st, style="label2" if on else "text", size=size, align="center",
                 where=f"flow_before_after[{i}]")
            if on:
                tag(slide, rx + rw - 40, ry - 10, "変更", style="alert", size=10, h=18)
            if i < len(steps) - 1:
                _arrow_line(slide, rx + rw / 2, ry + rh + 1, rx + rw / 2, ry + rh + 17, COLOR_GRAY_DARK, 1.5)
    _triangle(slide, a[0] + a[2] + 16, y + h / 2 - 30, 28, 60, fill=COLOR_PRIMARY_MID)


_ROMAN = ["i.", "ii.", "iii.", "iv.", "v.", "vi.", "vii.", "viii.", "ix.", "x."]
_IROHA = ["イ)", "ロ)", "ハ)", "ニ)", "ホ)", "ヘ)", "ト)", "チ)"]
_CIRC = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫"


@part("テキスト", "仕様書型の本文（①見出し＋段落、i./ii. と イ)ロ) の階層、__下線__で強調）")
def spec_text(slide, area, sections, size=13, numbered=True):
    """sections: [(見出し, 本文)]  本文は str（段落）か list（(1,"..") で i. ii.、(2,"..") で イ) ロ)）"""
    x, y, w, h = area
    blocks = []
    for k, (head, body) in enumerate(sections):
        ht = f"{_CIRC[k]} {head}" if numbered else head
        paras = []
        cnt = {1: 0, 2: 0}
        for lv, t in _norm_items(body if isinstance(body, list) else [body]):
            if lv == 1:
                cnt[1] += 1
                cnt[2] = 0
                t = f"{_ROMAN[cnt[1] - 1]} {t}"
            elif lv == 2:
                cnt[2] += 1
                t = f"{_IROHA[cnt[2] - 1]} {t}"
            paras.append((lv, t))
        hh = estimate_height(ht, w, size + 1) + 6
        bh = sum(estimate_height(t, w - 16 - lv * 16, size) for lv, t in paras) + 4 * len(paras) + 6
        blocks.append((ht, paras, hh, bh))
    total = sum(b[2] + b[3] + 8 for b in blocks)
    if total > h + 0.5:
        raise ValueError(f"[spec_text] 本文が領域に収まりません（推定 {total:.0f}pt > {h:.0f}pt）。文章を削るかスライドを分けてください。")
    cy = y
    for ht, paras, hh, bh in blocks:
        _textbox(slide, x, cy, w, hh, ht, size=size + 1, bold=True, color=COLOR_PRIMARY, where="spec_text見出し")
        cy += hh
        tb = slide.shapes.add_textbox(Pt(x + 16), Pt(cy), Pt(w - 16), Pt(bh))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Pt(3)
        tf.margin_top = tf.margin_bottom = Pt(2)
        for i, (lv, t) in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if i:
                p.space_before = Pt(4)
            if lv:
                pPr = p._p.get_or_add_pPr()
                ind = text_width(t.split(" ")[0] + " ", size)
                pPr.set("marL", str(int(Pt(lv * 16 + ind))))
                pPr.set("indent", str(-int(Pt(ind))))
            _add_runs(p, t, size, COLOR_TEXT, False)
        cy += bh + 8


@part("図解", "論点マップ（全論点を並べ、いま説明している論点だけを強調）")
def issue_map(slide, area, items, current, overview=None, size=12):
    """items: 論点の list（(番号, 文) でも可）。current: 強調する index。overview: 左に置く調査概要（任意）"""
    x, y, w, h = area
    if overview:
        a, b = split_cols(area, [1, 1.6], 20)
        titled_box(slide, a, "調査概要", overview, style="gray", size=size)
        x, y, w, h = b
    for i, (it, (rx, ry, rw, rh)) in enumerate(zip(items, split_rows((x, y, w, h), [1] * len(items), 6))):
        on = i == current
        t = f"{it[0]}. {it[1]}" if isinstance(it, tuple) else f"{i + 1}. {it}"
        sp = _rect(slide, rx, ry, rw, rh, fill=COLOR_PRIMARY if on else COLOR_WHITE,
                   line=None if on else COLOR_GRAY_LIGHT)
        _put_text(sp, t, size=size, bold=on, color=COLOR_WHITE if on else COLOR_GRAY_DARK, anchor="middle",
                  inset=(12, 2), where=f"issue_map[{i}]")


_MARKC = {"★": COLOR_PRIMARY, "✓": COLOR_PRIMARY_MID, "○": COLOR_PRIMARY_MID, "△": COLOR_GRAY_DARK,
          "◎": COLOR_PRIMARY, "×": COLOR_ACCENT}


@part("表", "★/✓ マトリクス（項目×企業など。列グループ見出し・凡例付き）")
def mark_matrix(slide, area, row_labels, col_labels, marks, groups=None, legend=None, size=12):
    """marks[行][列] = "★" / "✓" / "○" / "" など。groups: [(グループ名, 列数)] で上段に見出し。legend: {"★": "注力", ...}"""
    x, y, w, h = area
    if legend:
        lx = x + w
        for m, lab in reversed(list(legend.items())):
            tw = text_width(lab, 10) + 22
            lx -= tw + 6
            _textbox(slide, lx, y, 16, 18, m, size=12, bold=True, color=_MARKC.get(m, COLOR_TEXT), dense=True,
                     check=False, align="center", anchor="middle")
            _textbox(slide, lx + 16, y, tw - 16, 18, lab, size=10, dense=True, anchor="middle", check=False)
        y, h = y + 22, h - 22
    data = []
    if groups:
        row = [""]
        for g, n in groups:
            row += [g] + [""] * (n - 1)
        data.append(row)
    data.append([""] + list(col_labels))
    data += [[lab] + list(r) for lab, r in zip(row_labels, marks)]
    hdr = 2 if groups else 1

    def st(i, j, c):
        if i < hdr:
            return (COLOR_PRIMARY if i == 0 else COLOR_PRIMARY_MID) if groups else COLOR_PRIMARY, COLOR_WHITE, True, "center"
        if j == 0:
            return COLOR_GRAY_PALE, COLOR_TEXT, True, "left"
        return COLOR_WHITE, _MARKC.get((c or "")[:1], COLOR_TEXT), True, "center"
    gf = _native_table(slide, (x, y, w, h), data, [2] + [1] * len(col_labels), st, size, where="mark_matrix")
    if groups:
        tbl = gf.table
        c = 1
        for g, n in groups:
            if n > 1:
                tbl.cell(0, c).merge(tbl.cell(0, c + n - 1))
            c += n
    return gf


@part("表", "表＋行内の横棒（数値列を棒で見せる。国別比較など）")
def bar_table(slide, area, header, rows, bar_col, unit="", fmt="{:,.0f}", col_ratios=None, size=12):
    """rows の bar_col 列は数値。その列のセルに横棒と値を描く"""
    x, y, w, h = area
    vals = [r[bar_col] for r in rows]
    vmax = max(vals) or 1
    data = [list(header)] + [[("" if j == bar_col else c) for j, c in enumerate(r)] for r in rows]

    def st(i, j, c):
        if i == 0:
            return COLOR_PRIMARY, COLOR_WHITE, True, "center"
        return COLOR_WHITE, COLOR_TEXT, j == 0, "left"
    ratios = col_ratios or [1] * len(header)
    gf = _native_table(slide, area, data, ratios, st, size, where="bar_table")
    tbl = gf.table
    widths = _split(w, ratios, 0)
    bx = x + sum(widths[:bar_col])
    bw = widths[bar_col]
    cy = y + tbl.rows[0].height / 12700
    for i, v in enumerate(vals):
        rh = tbl.rows[i + 1].height / 12700
        lw = text_width(fmt.format(v) + unit, size) + 10
        maxw = bw - lw - 12
        _rect(slide, bx + 6, cy + rh * 0.25, max(maxw * v / vmax, 1), rh * 0.5,
              fill=COLOR_PRIMARY_MID if v == vmax else COLOR_PRIMARY_LIGHT)
        _textbox(slide, bx + 8 + maxw * v / vmax, cy, lw, rh, fmt.format(v) + unit, size=size, anchor="middle",
                 dense=True, where="bar_table値", check=False)
        cy += rh
    return gf


@part("構造", "件数付きの分類ツリー（全体 → 類型ごとの件数）")
def count_tree(slide, area, root, unit="件", size=12):
    """root = (ラベル, 件数, [子, ...])  子も同じ形。4階層まで"""
    def conv(n):
        lab, cnt, kids = n if len(n) == 3 else (n[0], n[1], [])
        t = f"{lab}\n**{cnt:,}{unit}**"
        return (t, [conv(k) for k in kids]) if kids else t
    logic_tree(slide, area, conv(root), size=size)


@part("数値", "絞り込みの棒（407件 → 293件 → 3件 のように段階的に減る過程）")
def reduction_bars(slide, area, steps, unit="件", size=12):
    """steps: [(ラベル, 値, 絞り込み条件 or None)]  条件は次の棒への矢印に添える"""
    x, y, w, h = area
    n = len(steps)
    vmax = max(s[1] for s in steps) or 1
    cols = split_cols(area, [1] * n, 96)
    base = y + h - 26
    top = y + 24
    fills = [COLOR_GRAY, COLOR_PRIMARY_LIGHT, COLOR_PRIMARY_MID, COLOR_PRIMARY]
    for i, ((lab, v, cond), (cx, _, cw, _)) in enumerate(zip(steps, cols)):
        bh = max((base - top) * v / vmax, 3)
        bw = min(cw * 0.6, 90)
        bx = cx + (cw - bw) / 2
        _rect(slide, bx, base - bh, bw, bh, fill=fills[min(i * 3 // max(n - 1, 1), 3)] if n > 1 else COLOR_PRIMARY)
        _textbox(slide, cx, base - bh - 24, cw, 22, f"{v:,}{unit}", size=size + 2, bold=True, color=COLOR_PRIMARY,
                 align="center", where="reduction_bars値")
        _textbox(slide, cx - 10, base + 2, cw + 20, 22, lab, size=size, align="center", where=f"reduction_bars[{i}]")
        if i < n - 1 and cond is not None:
            ax = cx + cw + 4
            _arrow_line(slide, ax + 4, base - (base - top) * 0.35, ax + 84, base - (base - top) * 0.35, COLOR_GRAY_DARK, 1.5)
            if cond:
                _textbox(slide, ax - 6, base - (base - top) * 0.35 - 48, 100, 44, cond, size=10, align="center",
                         anchor="bottom", color=COLOR_TEXT_SUB, dense=True, where="reduction_bars条件")


@part("カード", "事例紹介カード（画像・概要・ポイント・出所）")
def case_card(slide, area, title, overview, points, image_path=None, label=None, src=None, size=12):
    x, y, w, h = area
    _rect(slide, x, y, w, h, fill=COLOR_WHITE, line=COLOR_GRAY)
    iw = w * 0.36
    ia = (x + 10, y + 10, iw, h - 20)
    if image_path:
        from pptx_helpers import image
        _rect(slide, *ia, fill=COLOR_GRAY_PALE)
        image(slide, ia, image_path)
    else:
        from parts_media import image_placeholder
        image_placeholder(slide, ia, "画像")
    tx, tw = x + iw + 24, w - iw - 34
    cy = y + 10
    if label:
        tag(slide, tx, cy, label, style="mid")
        cy += 26
    _textbox(slide, tx, cy, tw, 28, title, size=size + 3, bold=True, color=COLOR_PRIMARY, where="case_card見出し")
    cy += 30
    oh = estimate_height(overview, tw - 6, size) + 8
    _textbox(slide, tx, cy, tw, oh, overview, size=size, where="case_card概要")
    cy += oh + 6
    _line(slide, tx, cy, tx + tw, cy, COLOR_GRAY_LIGHT, LINE_W)
    sh = 18 if src else 0
    _textbox(slide, tx, cy + 6, tw, y + h - cy - 16 - sh, points, size=size, bullets=True, where="case_cardポイント")
    if src:
        _textbox(slide, tx, y + h - 26, tw, 18, f"出所：{src}", size=10, color=COLOR_TEXT_SUB, dense=True,
                 where="case_card出所")


@part("テキスト", "コード・データ例のブロック（等幅フォント。JSON・設定例など）")
def code_block(slide, area, code, title=None, size=11):
    x, y, w, h = area
    if title:
        _box(slide, x, y, w, 24, title, style="label", size=size, dense=True)
        y, h = y + 24, h - 24
    _rect(slide, x, y, w, h, fill=COLOR_GRAY_PALE, line=COLOR_GRAY)
    lines = code.rstrip("\n").split("\n")
    need = len(lines) * size * 1.2 + 12
    longest = max(len(l) + sum(1 for c in l if ord(c) > 0x2E80) for l in lines) * 0.6 * size
    if need > h or longest > w - 16:
        raise ValueError(f"[code_block] コードが枠に収まりません（{len(lines)}行・最長 {longest:.0f}pt）。行を省略してください。")
    tb = slide.shapes.add_textbox(Pt(x + 6), Pt(y + 4), Pt(w - 12), Pt(h - 8))
    tf = tb.text_frame
    tf.word_wrap = False
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = l if l else " "
        _font(r.font, size, COLOR_TEXT, False)
        r.font.name = "Consolas"
    return tb


@part("レイアウト", "セクションタブ（本文の先頭に「4.1 システムの基本機能」の青タブ）。残りの本文領域を返す")
def section_tab(slide, label, area=None, size=13):
    x, y, w, h = area or BODY
    tw = text_width(label, size) * 1.1 + 34
    sp = _rect(slide, x, y, tw, 26, fill=COLOR_PRIMARY_MID, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    sp.adjustments[0] = 0.3
    _put_text(sp, label, size=size, color=COLOR_WHITE, bold=True, anchor="middle", inset=(14, 0), where="section_tab")
    _line(slide, x, y + 26, x + w, y + 26, COLOR_PRIMARY_MID, 1.5)
    return (x, y + 34, w, h - 34)
