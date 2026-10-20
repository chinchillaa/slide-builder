# -*- coding: utf-8 -*-
"""
pptx_helpers.py - 承認済みデザインを固定したスライド部品集

  * 座標・サイズの単位はすべて pt（スライドは 960 x 540 pt = 16:9）。
    領域は常に area = (x, y, w, h) のタプルで受け渡す。
  * 色・フォント・罫線・余白は本モジュールの定数で固定されている。
    関数が引数で受け付けない項目は「意図的に固定」されている。呼び出し側で上書きしないこと。
  * 本文フォントが下限未満、またはテキストが枠に収まらない場合は ValueError で停止する。
    → 文字数を削る / 枠を広げる / スライドを分ける で対処する（フォントを下げる抜け道は無い）。
  * テキスト内マークアップ: **太字**、!!強調（赤太字）!!、改行は "\n"。
  * 箇条書き items: ["項目", (1, "サブ項目"), ...]  タプル (level, text) で階層を指定。
"""
import math
import re
import unicodedata
from datetime import date, timedelta

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt

# ---------------------------------------------------------------- 寸法
SLIDE_W, SLIDE_H = 960, 540
MARGIN_X = 34
CONTENT_W = SLIDE_W - MARGIN_X * 2          # 892
BODY_TOP, BODY_BOTTOM = 118, 496
BODY = (MARGIN_X, BODY_TOP, CONTENT_W, BODY_BOTTOM - BODY_TOP)   # 本文領域 (34, 118, 892, 378)
GAP = 12                                    # 要素間の標準間隔

# ---------------------------------------------------------------- 文字
FONT = "Meiryo UI"
SZ_BREADCRUMB = 14
SZ_TITLE = 24
SZ_LEAD = 16
SZ_BODY = 14          # 本文の既定
SZ_HEADER = 14        # 列見出し
SZ_SOURCE = 9         # 出所・注記（chrome 扱い）
SZ_PAGE = 10.5
MIN_BODY_PT = 12      # 本文の下限。これ未満は ValueError
MIN_DENSE_PT = 10     # 表・ガント・グラフ・注記など高密度部品の下限
LINE_H = 1.25         # 行送り（推定用）

# ---------------------------------------------------------------- 色（hex 直書き禁止。必ず定数を使う）
COLOR_PRIMARY = "001F5F"        # 紺: ラベル塗り・タイトル罫線・強調ボックス
COLOR_PRIMARY_MID = "276CCC"    # 青: 二次ラベル・ハイライト
COLOR_PRIMARY_LIGHT = "8EB4E3"  # 薄青: グラフ二次系列・比較対象
COLOR_PRIMARY_PALE = "D5E2FF"   # 淡青: 交互行・ナビ選択・補足ボックス
COLOR_ACCENT = "C00000"         # 赤: 強調文字のみ（!!...!!）。塗りには使わない
COLOR_TEXT = "000000"
COLOR_TEXT_SUB = "404040"
COLOR_GRAY_DARK = "7E7F7E"      # 枠線・罫線
COLOR_GRAY = "BEBFBE"           # 矢印・非アクティブ
COLOR_GRAY_LIGHT = "D8D9D8"     # 区切り線
COLOR_GRAY_PALE = "F1F2F1"      # 薄い背景
COLOR_WHITE = "FFFFFF"
COLOR_STATUS_OK = "3C8D5A"      # 信号の青（良好）: 状態表示の部品専用
COLOR_STATUS_WARN = "E0A100"    # 信号の黄（注意）: 状態表示の部品専用
COLOR_STATUS_NG = COLOR_ACCENT  # 信号の赤（問題）
SERIES_COLORS = [COLOR_PRIMARY, COLOR_PRIMARY_MID, COLOR_PRIMARY_LIGHT,
                 COLOR_GRAY, COLOR_GRAY_DARK, COLOR_PRIMARY_PALE]
PALETTE = {COLOR_PRIMARY, COLOR_PRIMARY_MID, COLOR_PRIMARY_LIGHT, COLOR_PRIMARY_PALE,
           COLOR_ACCENT, COLOR_TEXT, COLOR_TEXT_SUB, COLOR_GRAY_DARK, COLOR_GRAY,
           COLOR_GRAY_LIGHT, COLOR_GRAY_PALE, COLOR_WHITE, COLOR_STATUS_OK, COLOR_STATUS_WARN}

LINE_W = 0.75         # 標準罫線
TITLE_RULE_W = 2.25   # タイトル下罫線

_STATE = {"footer": "", "cover": False}
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_MARK = re.compile(r"(\*\*.+?\*\*|!!.+?!!|__.+?__)")


# ================================================================ 内部ユーティリティ
def _rgb(hex_):
    return RGBColor.from_string(hex_)


def _check_size(size, dense, where):
    lo = MIN_DENSE_PT if dense else MIN_BODY_PT
    if size < lo:
        raise ValueError(
            f"[{where}] フォントサイズ {size}pt は下限 {lo}pt 未満です。"
            f"文字数を減らす・枠を広げる・スライドを分けるで対処してください。")


def _char_w(ch):
    return 1.0 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 0.6


def _plain(text):
    return text.replace("**", "").replace("!!", "").replace("__", "")


def text_width(text, size):
    """1行テキストの推定幅(pt)"""
    return max((sum(_char_w(c) for c in seg) for seg in _plain(text).split("\n")), default=0) * size


_TOKEN = re.compile(r"[A-Za-z0-9][A-Za-z0-9.,:%/_\-]*")


def check_tokens(text, width, size, where):
    """英単語・数値が途中で折り返されないか検査（日本語と違い単語の途中改行は読みにくい）"""
    for tok in _TOKEN.findall(_plain(text)):
        if len(tok) > 1 and len(tok) * 0.6 * size > width:
            raise ValueError(f"[{where}] 「{tok}」が幅 {width:.0f}pt に収まらず単語の途中で改行されます"
                             f"（{size}pt）。枠を広げるか表記を短くしてください。")


def fit_one_line(text, width, size):
    """1行に収まる最大の文字サイズ（size を上限）"""
    return min(size, width / max(text_width(text, 1) * 1.15, 0.1))   # 太字の数字は推定より広いので余裕を持たせる


def estimate_lines(text, width, size):
    """折り返しを考慮した推定行数"""
    n = 0
    for seg in _plain(text).split("\n"):
        w = sum(_char_w(c) for c in seg) * size * 1.03
        n += max(1, math.ceil(w / max(width, 1) - 1e-6))
    return n


def _norm_items(items):
    if isinstance(items, str):
        items = [items]
    out = []
    for it in items:
        if isinstance(it, tuple) and len(it) == 2 and isinstance(it[0], int):
            out.append((it[0], it[1]))
        else:
            out.append((0, str(it)))
    return out


def _indent(level, size):
    return size * (1.0 + 1.1 * level), size * 1.0     # marL, hanging


def estimate_height(items, width, size, bullets=False):
    """items を幅 width に流し込んだときの推定高さ(pt)"""
    total = 0.0
    norm = _norm_items(items)
    for i, (lv, t) in enumerate(norm):
        w = width - (_indent(lv, size)[0] if bullets else 0)
        total += estimate_lines(t, w, size) * size * LINE_H
        if i:
            total += size * 0.3 if bullets else 0
    return total


def _set_ea(rPr):
    latin = rPr.find(qn("a:latin"))
    prev = latin
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            if prev is not None:
                prev.addnext(el)
            else:
                rPr.append(el)
        el.set("typeface", FONT)
        prev = el


def _font(font, size, color, bold):
    font.name = FONT
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = _rgb(color)
    _set_ea(font._rPr)


def _add_runs(p, text, size, color, bold):
    lines = text.split("\n")
    for li, line in enumerate(lines):
        if li:
            p.add_line_break()
        for part in _MARK.split(line):
            if not part:
                continue
            b, c, u = bold, color, False
            if len(part) >= 4 and part.startswith("__") and part.endswith("__"):
                part, u = part[2:-2], True
            elif len(part) >= 4 and part.startswith("**") and part.endswith("**"):
                part, b = part[2:-2], True
            elif len(part) >= 4 and part.startswith("!!") and part.endswith("!!"):
                part, b, c = part[2:-2], True, COLOR_ACCENT
            r = p.add_run()
            r.text = part
            _font(r.font, size, c, b)
            if u:
                r.font.underline = True


def _bullet(p, level, size):
    pPr = p._p.get_or_add_pPr()
    marL, hang = _indent(level, size)
    pPr.set("marL", str(int(Pt(marL))))
    pPr.set("indent", str(-int(Pt(hang))))
    for tag in ("a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))


def _fill_text(tf, items, size, color, bold, align, anchor, bullets, inset):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    ix, iy = inset
    tf.margin_left = tf.margin_right = Pt(ix)
    tf.margin_top = tf.margin_bottom = Pt(iy)
    tf.vertical_anchor = _ANCHOR[anchor]
    for i, (lv, t) in enumerate(_norm_items(items)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        if bullets and i:
            p.space_before = Pt(size * 0.3)
        if bullets:
            _bullet(p, lv, size)
        _add_runs(p, t, size, color, bold)


def _put_text(shape, items, *, size, color=COLOR_TEXT, bold=False, align="left",
              anchor="top", bullets=False, dense=False, inset=(6, 4), where="text", check=True,
              exempt=False):
    if not exempt:          # exempt はページ番号・出所など chrome 専用
        _check_size(size, dense, where)
    k = 1.0
    try:
        if shape.auto_shape_type in (MSO_SHAPE.OVAL, MSO_SHAPE.DIAMOND):
            k = 0.7 if shape.auto_shape_type == MSO_SHAPE.OVAL else 0.5   # 楕円・ひし形の文字領域は内接矩形
    except (AttributeError, ValueError, NotImplementedError):
        pass
    w = shape.width / 12700 * k - inset[0] * 2
    h = shape.height / 12700 * k - inset[1] * 2
    if check:
        for lv, t in _norm_items(items):
            check_tokens(t, w - (_indent(lv, size)[0] if bullets else 0), size, where)
        est = estimate_height(items, w, size, bullets)
        if est > h + 0.5:
            raise ValueError(
                f"[{where}] テキストが枠に収まりません: 推定高さ {est:.0f}pt > 枠 {h:.0f}pt"
                f"（{size}pt, 幅 {w:.0f}pt）。文字数を減らす・枠を広げる・スライドを分けるで対処してください。"
                f" 先頭: {_plain(_norm_items(items)[0][1])[:30]}")
    _fill_text(shape.text_frame, items, size, color, bold, align, anchor, bullets, inset)
    return shape


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=LINE_W, shape=MSO_SHAPE.RECTANGLE, name=None):
    sp = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    sp.shadow.inherit = False
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    if name:
        sp.name = name
    return sp


def _textbox(slide, x, y, w, h, items, *, name=None, **kw):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        tb.name = name
    kw.setdefault("inset", (3, 2))
    return _put_text(tb, items, **kw)


def _no_effect(shape):
    """テーマ由来の影を消す（コネクタは既定で影が付くため）"""
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        ln = spPr.find(qn("a:ln"))
        el = spPr.makeelement(qn("a:effectLst"), {})
        ln.addnext(el) if ln is not None else spPr.append(el)


def _line(slide, x1, y1, x2, y2, color=COLOR_GRAY_DARK, width=LINE_W, dash=False, name=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if name:
        ln.name = name
    _no_effect(ln)
    return ln


def _triangle(slide, x, y, w, h, direction="right", fill=COLOR_GRAY_LIGHT):
    pts = {"right": [(x, y), (x + w, y + h / 2), (x, y + h)],
           "down": [(x, y), (x + w, y), (x + w / 2, y + h)]}[direction]
    fb = slide.shapes.build_freeform(Pt(pts[0][0]), Pt(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Pt(px), Pt(py)) for px, py in pts[1:]], close=True)
    sp = fb.convert_to_shape()
    sp.shadow.inherit = False
    sp.fill.solid()
    sp.fill.fore_color.rgb = _rgb(fill)
    sp.line.fill.background()
    sp.name = "arrow"
    return sp


def _apply_theme_fonts(prs):
    """テーマの和文・欧文フォントを Meiryo UI にする（手で追加したテキストや表・グラフにも効く）"""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for tag in ("majorFont", "minorFont"):
        f = root.find(f".//a:fontScheme/a:{tag}", ns)
        for sub in ("latin", "ea"):
            el = f.find(f"a:{sub}", ns)
            if el is not None:
                el.set("typeface", FONT)
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if hasattr(part, "_element"):
        part._element = root
    else:
        part._blob = blob


def _split(total, ratios, gap):
    avail = total - gap * (len(ratios) - 1)
    s = float(sum(ratios))
    return [avail * r / s for r in ratios]


# ================================================================ 部品レジストリ（一覧は list_parts.py で生成）
REGISTRY = {}


def part(category, desc):
    """部品として登録するデコレータ"""
    def deco(fn):
        REGISTRY[fn.__name__] = {"category": category, "desc": desc, "fn": fn}
        return fn
    return deco


# ================================================================ デッキ・枠（chrome）
def new_deck(footer=""):
    """16:9 の空デッキを作る。footer はスライド右下の小さな表記（例: "©2026 ○○部"）"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(SLIDE_W), Pt(SLIDE_H)
    _apply_theme_fonts(prs)
    _STATE["footer"] = footer
    _STATE["cover"] = False
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _chrome_footer(prs, slide, dark=False):
    n = len(prs.slides) - (1 if _STATE["cover"] else 0)   # 表紙は番号なし・次から 1
    col = COLOR_PRIMARY_PALE if dark else COLOR_TEXT_SUB
    _textbox(slide, 430, 514, 100, 18, str(n), size=SZ_PAGE, color=col, align="center",
             dense=True, name="chrome:page", check=False)
    if _STATE["footer"]:
        _textbox(slide, 640, 516, 300, 16, _STATE["footer"], size=8, color=col, align="right",
                 dense=True, name="chrome:footer", check=False, exempt=True)


@part("スライド枠", "表紙")
def cover_slide(prs, title, subtitle="", date_text="", org="", recipient=""):
    """表紙。title は最大2行（28pt）"""
    s = _blank(prs)
    _STATE["cover"] = True
    if recipient:
        _textbox(s, 50, 40, 600, 24, recipient, size=14, name="chrome:recipient")
    _textbox(s, 50, 140, 860, 110, title, size=28, bold=True, anchor="bottom",
             name="chrome:cover_title", where="表紙タイトル")
    _line(s, 50, 258, 910, 258, COLOR_PRIMARY, TITLE_RULE_W, name="chrome:rule")
    if subtitle:
        _textbox(s, 50, 266, 860, 36, subtitle, size=20, bold=True, color=COLOR_PRIMARY,
                 name="chrome:subtitle", where="表紙サブタイトル")
    meta = [t for t in (date_text, org) if t]
    if meta:
        _textbox(s, 50, 400, 600, 60, meta, size=16, name="chrome:meta")
    return s


@part("スライド枠", "章の区切り（紺の全面塗り）")
def section_slide(prs, title, items=None):
    """セクション区切り（紺ベタ）。items は配下の見出し一覧"""
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, fill=COLOR_PRIMARY, name="chrome:bg")
    _textbox(s, 70, 190, 820, 56, title, size=28, bold=True, color=COLOR_WHITE, anchor="bottom",
             name="chrome:section_title", where="セクションタイトル")
    _line(s, 70, 252, 890, 252, COLOR_WHITE, 1.0, name="chrome:rule")
    if items:
        _textbox(s, 90, 262, 800, 220, items, size=16, color=COLOR_WHITE,
                 name="chrome:section_items", where="セクション項目")
    _chrome_footer(prs, s, dark=True)
    return s


@part("スライド枠", "本文スライドの枠（パンくず・タイトル・罫線・リード文）")
def content_slide(prs, title, lead="", breadcrumb="", nav=None):
    """本文スライド。パンくず(14pt太字) / タイトル(24pt太字・1行) / 紺罫線 / リード文(16pt・2行以内)。
    nav=(tabs, active_index) で右上にナビタブ。本文は BODY 領域に置く。"""
    s = _blank(prs)
    title_w = CONTENT_W
    if nav:
        tabs, active = nav
        title_w -= nav_tabs(s, tabs, active) + GAP
    if breadcrumb:
        _textbox(s, MARGIN_X, 5, title_w, 22, breadcrumb, size=SZ_BREADCRUMB, bold=True,
                 name="chrome:breadcrumb", where="パンくず")
    _textbox(s, MARGIN_X, 22, title_w, 38, title, size=SZ_TITLE, bold=True, anchor="middle",
             name="chrome:title", where="タイトル（1行に収める）")
    _line(s, 0, 60, SLIDE_W, 60, COLOR_PRIMARY, TITLE_RULE_W, name="chrome:rule")
    if lead:
        _textbox(s, MARGIN_X, 62, CONTENT_W, 50, lead, size=SZ_LEAD,
                 name="chrome:lead", where="リード文（2行以内）")
    _chrome_footer(prs, s)
    return s


@part("スライド枠", "右上のナビタブ（現在地の表示）")
def nav_tabs(slide, tabs, active):
    """右上のナビタブ（現在地表示）。占有幅を返す"""
    tw, th = 100, 24
    x0 = SLIDE_W - 4 - tw * len(tabs)
    for i, t in enumerate(tabs):
        on = i == active
        sp = _rect(slide, x0 + i * tw, 5, tw, th, fill=COLOR_PRIMARY_PALE if on else COLOR_WHITE,
                   line=COLOR_PRIMARY if on else COLOR_GRAY, line_w=1.5 if on else LINE_W,
                   name="chrome:nav")
        _put_text(sp, t, size=9, bold=on, color=COLOR_PRIMARY_MID if on else COLOR_GRAY_DARK,
                  align="center", anchor="middle", dense=True, inset=(2, 1), where="ナビタブ", exempt=True)
    return tw * len(tabs) + 4


@part("スライド枠", "目次（2段組まで自動）")
def toc_slide(prs, items, title="目次"):
    """目次。items: [(level, text, page), ...]  level 0=章(太字) / 1=節。15行超で2段組"""
    s = content_slide(prs, title)
    row_h = 24
    per_col = 15
    if len(items) > per_col * 2:
        raise ValueError(f"[目次] 項目が {len(items)} 行あり 2段組（{per_col*2}行）を超えます。節を省略してください。")
    cols = 1 if len(items) <= per_col else 2
    cw = (CONTENT_W - 40) / cols
    for i, (lv, text, page) in enumerate(items):
        c, r = divmod(i, per_col)
        x = MARGIN_X + 40 + c * (cw + 20) + lv * 22
        y = 90 + r * row_h
        sz = 16 if lv == 0 else 14
        _textbox(s, x, y, cw - 90 - lv * 22, row_h, text, size=sz, bold=lv == 0, anchor="middle",
                 where="目次")
        _textbox(s, MARGIN_X + 40 + c * (cw + 20) + cw - 90, y, 70, row_h, str(page), size=sz - 2,
                 color=COLOR_TEXT_SUB, align="right", anchor="middle", where="目次頁", dense=True)
    return s


@part("補助", "左下の出所・注記")
def source(slide, text):
    """左下の出所・注記（9pt）"""
    if not text.startswith(("出所", "出典", "注", "※")):
        text = "出所：" + text
    return _textbox(slide, MARGIN_X, 497, CONTENT_W - 180, 16, text, size=SZ_SOURCE,
                    color=COLOR_TEXT_SUB, dense=True, name="chrome:source", where="出所", exempt=True)


# ================================================================ 本文部品
@part("レイアウト", "列見出し（下線付き）")
def column_header(slide, x, y, w, text, size=SZ_HEADER):
    """列見出し（中央寄せテキスト + 下線）。消費した高さを返す"""
    _textbox(slide, x, y, w, 26, text, size=size, align="center", anchor="bottom", where="列見出し")
    _line(slide, x, y + 30, x + w, y + 30, COLOR_TEXT, LINE_W)
    return 38


@part("レイアウト", "左右分割（見出し・列間矢印は任意）")
def split_columns(slide, area, ratios, titles=None, arrow=False):
    """領域を左右に分割し、各列の本文領域を返す。
    titles で列見出し、arrow=True で列間にグレー三角矢印。
    戻り値: (areas, arrow_areas)  ※見出し位置は全列で必ず揃う"""
    x, y, w, h = area
    gap = 34 if arrow else GAP
    widths = _split(w, ratios, gap)
    top = y
    if titles:
        if len(titles) != len(ratios):
            raise ValueError("titles と ratios の数が一致しません")
        cx = x
        for cw, t in zip(widths, titles):
            used = column_header(slide, cx, y, cw, t) if t else 38
            cx += cw + gap
        top = y + used
    areas, arrows = [], []
    cx = x
    for i, cw in enumerate(widths):
        areas.append((cx, top, cw, y + h - top))
        if i < len(widths) - 1:
            ax = cx + cw + 8
            arrows.append((ax, top, gap - 16, y + h - top))
            if arrow:
                ah = min((y + h - top) * 0.5, 140)
                _triangle(slide, ax, top + (y + h - top - ah) / 2, gap - 16, ah)
        cx += cw + gap
    return areas, arrows


@part("テキスト", "箇条書き（枠・小見出しは任意）")
def bullets(slide, area, items, size=SZ_BODY, title=None, box=False, anchor="top"):
    """箇条書き。title で紺太字の小見出し、box=True でグレー枠"""
    x, y, w, h = area
    if box:
        _rect(slide, x, y, w, h, fill=COLOR_WHITE, line=COLOR_GRAY_DARK)
        x, y, w, h = x + 6, y + 6, w - 12, h - 12
    if title:
        th = size * 1.7
        _textbox(slide, x, y, w, th, title, size=size, bold=True, color=COLOR_PRIMARY, where="小見出し")
        y, h = y + th, h - th
    return _textbox(slide, x, y, w, h, items, size=size, bullets=True, anchor=anchor, where="箇条書き")


@part("テキスト", "地の文")
def text(slide, area, body, size=SZ_BODY, bold=False, align="left", anchor="top", color=COLOR_TEXT):
    """箇条書きにしない地の文"""
    x, y, w, h = area
    return _textbox(slide, x, y, w, h, body, size=size, bold=bold, align=align, anchor=anchor,
                    color=color, where="本文")


_CELL = {
    # style: (fill, line, text color, bold, align, bullets)
    "label":  (COLOR_PRIMARY, None, COLOR_WHITE, True, "center", False),
    "label2": (COLOR_PRIMARY_MID, None, COLOR_WHITE, True, "center", False),
    "pale":   (COLOR_PRIMARY_PALE, None, COLOR_PRIMARY, True, "center", False),
    "gray":   (COLOR_GRAY_PALE, COLOR_GRAY_DARK, COLOR_TEXT, False, "center", False),
    "text":   (COLOR_WHITE, COLOR_GRAY_DARK, COLOR_TEXT, False, "left", True),
    "plain":  (None, None, COLOR_TEXT, False, "left", True),
}


@part("表", "カード型の表（紺ラベル＋本文、結合・矢印あり）")
def card_table(slide, area, rows, col_ratios, styles, headers=None, arrow_before=(),
               row_ratios=None, size=SZ_BODY):
    """カード型の表（ラベル列 + 本文列）。資料で最も多用する部品。
    rows: [[cell, ...], ...]  cell は str / list(箇条書き) / None(上のセルと結合: label 系のみ)
    styles: 列ごとのスタイル "label"(紺) "label2"(青) "pale"(淡青) "gray" "text"(枠+箇条書き) "plain"
    headers: 列見出し（None で省略、要素に "" で個別省略）
    arrow_before: この列番号の手前にグレー三角矢印を置く（例: (3,)）"""
    x, y, w, h = area
    n = len(col_ratios)
    if len(styles) != n or any(len(r) != n for r in rows):
        raise ValueError("[card_table] 列数が col_ratios / styles / rows で一致しません")
    gap, arrow_gap = 6, 30
    total_gap = gap * (n - 1) + (arrow_gap - gap) * len(arrow_before)
    s = float(sum(col_ratios))
    widths = [(w - total_gap) * r / s for r in col_ratios]
    xs, cx = [], x
    for j, cw in enumerate(widths):
        if j in arrow_before:
            cx += arrow_gap - gap
        xs.append(cx)
        cx += cw + gap
    top = y
    if headers:
        for j, t in enumerate(headers):
            if t:
                column_header(slide, xs[j], y, widths[j], t)
        top = y + 38
    rh_list = _split(y + h - top, row_ratios or [1] * len(rows), gap)
    ys, cy = [], top
    for rh in rh_list:
        ys.append(cy)
        cy += rh + gap
    for j in range(n):
        fill, line, col, bold, align, bul = _CELL[styles[j]]
        i = 0
        while i < len(rows):
            cell = rows[i][j]
            span = 1
            while i + span < len(rows) and rows[i + span][j] is None:
                span += 1
            if cell is None:
                i += 1
                continue
            hh = sum(rh_list[i:i + span]) + gap * (span - 1)
            sp = _rect(slide, xs[j], ys[i], widths[j], hh, fill=fill, line=line)
            is_list = isinstance(cell, list)
            _put_text(sp, cell, size=size, color=col, bold=bold, align="left" if is_list else align,
                      anchor="middle", bullets=bul and is_list, where=f"card_table[{i},{j}]")
            i += span
    for j in arrow_before:
        for i in range(len(rows)):
            _triangle(slide, xs[j] - arrow_gap + 6, ys[i] + 4, arrow_gap - 12, rh_list[i] - 8)
    return xs, widths, ys, rh_list


def _cell_border(cell, color, width):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.makeelement(qn(tag), {"w": str(int(Pt(width))), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(sf)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        tcPr.append(ln)


@part("表", "ネイティブの表（データの多い一覧）")
def make_table(slide, area, header, rows, col_ratios=None, size=12, align=None,
               first_col_label=False, zebra=False, stretch=False):
    """PowerPoint ネイティブの表（編集可能）。データ量の多い一覧向け。
    align: 列ごと "l"/"c"/"r"。cell は str か list(箇条書き)。
    stretch=True で行高を広げて領域を埋める。収まらなければ ValueError"""
    x, y, w, h = area
    _check_size(size, True, "make_table")
    nc = len(header)
    col_ratios = col_ratios or [1] * nc
    widths = _split(w, col_ratios, 0)
    align = align or ["l"] * nc
    amap = {"l": "left", "c": "center", "r": "right"}
    mx, my = 5, 3
    all_rows = [header] + rows
    heights = []
    for r in all_rows:
        need = max(estimate_height(c if c is not None else "", widths[j] - mx * 2, size,
                                   isinstance(c, list)) for j, c in enumerate(r))
        heights.append(max(need + my * 2 + 2, size * 2.0))
    total = sum(heights)
    if total > h + 0.5:
        raise ValueError(f"[make_table] 表が領域に収まりません: 推定 {total:.0f}pt > {h:.0f}pt。"
                         f"行・文字を減らすか、表を分割してください。")
    if stretch:
        extra = (h - total) / len(heights)
        heights = [hh + extra for hh in heights]
    gf = slide.shapes.add_table(len(all_rows), nc, Pt(x), Pt(y), Pt(w), Pt(sum(heights)))
    tbl = gf.table
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # No Style, No Grid
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Pt(cw)
    for i, hh in enumerate(heights):
        tbl.rows[i].height = Pt(hh)
    for i, r in enumerate(all_rows):
        for j, c in enumerate(r):
            cell = tbl.cell(i, j)
            _cell_border(cell, COLOR_GRAY_DARK, 0.5)
            if i == 0:
                fill, col, bold, al = COLOR_PRIMARY, COLOR_WHITE, True, "center"
            elif first_col_label and j == 0:
                fill, col, bold, al = COLOR_PRIMARY_PALE, COLOR_PRIMARY, True, "left"
            else:
                fill = COLOR_GRAY_PALE if zebra and i % 2 == 0 else COLOR_WHITE
                col, bold, al = COLOR_TEXT, False, amap[align[j]]
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            cell.margin_left = cell.margin_right = Pt(mx)
            cell.margin_top = cell.margin_bottom = Pt(my)
            tf = cell.text_frame
            tf.word_wrap = True
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for k, (lv, t) in enumerate(_norm_items(c if c is not None else "")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.alignment = _ALIGN[al]
                if isinstance(c, list):
                    _bullet(p, lv, size)
                _add_runs(p, t, size, col, bold)
    return gf


@part("フロー", "シェブロン見出し＋説明枠のステップ")
def step_flow(slide, area, steps, size=SZ_BODY, active=None, header_h=40):
    """横向きのステップ（シェブロン見出し + 各ステップの説明枠）。
    steps: [(見出し, items), ...]  active を指定するとそれ以外の見出しをグレーにする"""
    x, y, w, h = area
    n = len(steps)
    widths = _split(w, [1] * n, 8)
    cx = x
    for i, (title, items) in enumerate(steps):
        on = active is None or i == active
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        sp = _rect(slide, cx, y, widths[i] + (6 if i < n - 1 else 0), header_h,
                   fill=COLOR_PRIMARY if on else COLOR_GRAY, shape=shp)
        _put_text(sp, title, size=size, color=COLOR_WHITE, bold=True, align="center", anchor="middle",
                  inset=(18 if i else 8, 2), where=f"step_flow見出し[{i}]")
        if items:
            box = _rect(slide, cx, y + header_h + 8, widths[i], h - header_h - 8,
                        fill=COLOR_WHITE, line=COLOR_PRIMARY if on else COLOR_GRAY_DARK)
            _put_text(box, items, size=size, bullets=True, inset=(8, 8), where=f"step_flow本文[{i}]")
        cx += widths[i] + 8
    return widths


@part("テキスト", "結論・キーメッセージの強調ボックス")
def message_box(slide, area, body, style="primary", size=16):
    """結論・キーメッセージの強調ボックス。style: primary(紺ベタ) / pale(淡青) / outline(紺枠)"""
    x, y, w, h = area
    fill, line, col = {"primary": (COLOR_PRIMARY, None, COLOR_WHITE),
                       "pale": (COLOR_PRIMARY_PALE, None, COLOR_PRIMARY),
                       "outline": (COLOR_WHITE, COLOR_PRIMARY, COLOR_PRIMARY)}[style]
    sp = _rect(slide, x, y, w, h, fill=fill, line=line, line_w=1.5)
    return _put_text(sp, body, size=size, color=col, bold=True, align="center", anchor="middle",
                     inset=(10, 6), where="message_box")


@part("補助", "グレーの三角矢印")
def arrow(slide, area, direction="right"):
    """グレーの三角矢印。direction: right / down"""
    x, y, w, h = area
    return _triangle(slide, x, y, w, h, direction)


@part("数値", "大きな数字のKPIカード")
def kpi(slide, area, value, label, note=None, unit=""):
    """大きな数値で示す KPI カード"""
    x, y, w, h = area
    _rect(slide, x, y, w, h, fill=COLOR_WHITE, line=COLOR_GRAY_DARK)
    _textbox(slide, x + 6, y + 6, w - 12, 24, label, size=14, bold=True, align="center", where="kpiラベル")
    vh = h - 36 - (34 if note else 6)
    vs = fit_one_line(f"{value}{unit}", w - 16, min(40, vh / 1.4))
    _check_size(vs, False, "kpi値")
    _textbox(slide, x + 6, y + 32, w - 12, vh, f"{value}{unit}", size=vs, bold=True,
             color=COLOR_PRIMARY, align="center", anchor="middle", where="kpi値")
    if note:
        _textbox(slide, x + 6, y + h - 36, w - 12, 30, note, size=12, align="center",
                 color=COLOR_TEXT_SUB, where="kpi注記")


@part("画像", "画像を縦横比を保って配置")
def image(slide, area, path):
    """画像を領域内に縦横比を保って中央配置"""
    from PIL import Image
    x, y, w, h = area
    with Image.open(path) as im:
        iw, ih = im.size
    sc = min(w / iw, h / ih)
    dw, dh = iw * sc, ih * sc
    return slide.shapes.add_picture(path, Pt(x + (w - dw) / 2), Pt(y + (h - dh) / 2), Pt(dw), Pt(dh))


_CHART = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
          "bar": XL_CHART_TYPE.BAR_CLUSTERED, "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
          "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
          "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
          "area_stacked": XL_CHART_TYPE.AREA_STACKED, "column_100": XL_CHART_TYPE.COLUMN_STACKED_100,
          "bar_100": XL_CHART_TYPE.BAR_STACKED_100, "radar": XL_CHART_TYPE.RADAR_MARKERS}


@part("グラフ", "ネイティブグラフ（棒/積上/横棒/折れ線/面/円/ドーナツ/100%積上）")
def chart(slide, area, kind, categories, series, title=None, unit=None, number_format="#,##0",
          labels=None, legend=None, highlight=None, size=12):
    """PowerPoint ネイティブグラフ（編集可能）。
    kind: column / column_stacked / bar / bar_stacked / line / pie / doughnut
    series: {"系列名": [値, ...]}（順序 = 色順 SERIES_COLORS）
    highlight: 単一系列の棒で強調するカテゴリ index（他は薄青）"""
    x, y, w, h = area
    _check_size(size, True, "chart")
    if labels is None:   # 既定: 重なりやすい種類ではラベルを出さない
        labels = not (kind.startswith("area") or kind == "radar" or (kind == "line" and len(series) > 1))
    if title or unit:
        if title:
            _textbox(slide, x, y, w, 24, title, size=14, bold=True, align="center", where="グラフタイトル")
        if unit:
            _textbox(slide, x + w - 160, y + 22, 160, 16, f"（単位：{unit}）", size=10, align="right",
                     color=COLOR_TEXT_SUB, dense=True, where="グラフ単位")
        y, h = y + 40, h - 40
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series.items():
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(_CHART[kind], Pt(x), Pt(y), Pt(w), Pt(h), cd)
    ch = gf.chart
    ch.has_title = False          # タイトルは上のテキストで出す
    ch.font.name, ch.font.size = FONT, Pt(size)
    ch.font.color.rgb = _rgb(COLOR_TEXT)
    for rpr in ch._chartSpace.iter(qn("a:defRPr")):
        _set_ea(rpr)
    pie = kind in ("pie", "doughnut")
    show_legend = legend if legend is not None else (pie or len(series) > 1)
    ch.has_legend = show_legend
    if show_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    plot = ch.plots[0]
    if pie:
        pts = plot.series[0].points
        for i in range(len(categories)):
            pts[i].format.fill.solid()
            pts[i].format.fill.fore_color.rgb = _rgb(SERIES_COLORS[i % len(SERIES_COLORS)])
            pts[i].format.line.color.rgb = _rgb(COLOR_WHITE)
    else:
        if kind != "line":
            plot.gap_width = 60
            if "stacked" in kind or "100" in kind:
                plot.overlap = 100
        for i, s in enumerate(plot.series):
            col = SERIES_COLORS[i % len(SERIES_COLORS)]
            if kind in ("line", "radar"):
                s.format.line.color.rgb = _rgb(col)
                s.format.line.width = Pt(2.25)
                s.smooth = False
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = _rgb(col)
                s.marker.format.line.color.rgb = _rgb(col)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = _rgb(col)
        if highlight is not None and len(series) == 1 and kind not in ("line", "radar") and not kind.startswith("area"):
            s = plot.series[0]
            for i in range(len(categories)):
                p = s.points[i]
                p.format.fill.solid()
                p.format.fill.fore_color.rgb = _rgb(COLOR_PRIMARY if i == highlight else COLOR_PRIMARY_LIGHT)
        va, ca = ch.value_axis, ch.category_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _rgb(COLOR_GRAY_LIGHT)
        va.major_gridlines.format.line.width = Pt(0.5)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(size)
        va.tick_labels.number_format = "0%" if "100" in kind else number_format
        va.tick_labels.number_format_is_linked = False
        ca.format.line.color.rgb = _rgb(COLOR_GRAY_DARK)
        ca.tick_labels.font.size = Pt(size)
        ca.has_major_gridlines = False
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(size)
        dl.number_format = "0%" if pie else number_format
        dl.number_format_is_linked = False
        if pie:
            dl.show_percentage, dl.show_value = True, False
            dl.font.color.rgb = _rgb(COLOR_WHITE)
            dl.font.bold = True
        elif kind == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
        elif kind.startswith("area") or kind == "radar":
            pass
        elif "stacked" in kind or "100" in kind:
            dl.position = XL_LABEL_POSITION.CENTER
            dl.font.color.rgb = _rgb(COLOR_WHITE)
        else:
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    return gf


def _months(start, end):
    d = date(start.year, start.month, 1)
    while d <= end:
        nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
        yield max(d, start), min(nxt - timedelta(days=1), end)
        d = nxt


@part("スケジュール", "ガントチャート（月・週の目盛り、マイルストーン）")
def gantt(slide, area, start, end, rows, bars, milestones=(), label_w=150, size=12):
    """スケジュール（ガント）。
    rows: 行ラベル list。 bars: [(row, 開始date, 終了date, ラベル, style="primary"|"mid"|"pale", lane=0)]
    milestones: [(date, ラベル)]  週数が26以下なら週の目盛りも表示"""
    x, y, w, h = area
    _check_size(size, True, "gantt")
    gx, gw = x + label_w + 6, w - label_w - 6
    days = (end - start).days + 1
    def dx(d):
        return gx + (d - start).days / days * gw
    mh = 20
    _rect(slide, gx, y, gw, mh, fill=COLOR_PRIMARY)
    for ms, me in _months(start, end):
        mx0, mx1 = dx(ms), dx(me + timedelta(days=1))
        lab = f"{ms.month}月" if mx1 - mx0 > 30 else str(ms.month)
        _textbox(slide, mx0, y, mx1 - mx0, mh, lab, size=11, bold=True, color=COLOR_WHITE,
                 align="center", anchor="middle", dense=True, where="gantt月", inset=(1, 0),
                 check=mx1 - mx0 > 14)
    top = y + mh
    weeks = []
    d = start - timedelta(days=start.weekday())
    while d <= end:
        weeks.append(max(d, start))
        d += timedelta(days=7)
    if len(weeks) <= 26:
        wh = 18
        _rect(slide, gx, top, gw, wh, fill=COLOR_GRAY_PALE)
        for wd in weeks:
            wx = dx(wd)
            if gx + gw - wx > 18:
                _textbox(slide, wx, top, 34, wh, f"{wd.month}/{wd.day}", size=10, color=COLOR_TEXT_SUB,
                         dense=True, anchor="middle", inset=(1, 0), where="gantt週", check=False)
        top += wh
    bottom = y + h - (18 if milestones else 0)
    for ms, _ in _months(start, end):
        if ms != start:
            _line(slide, dx(ms), y + mh, dx(ms), bottom, COLOR_GRAY_LIGHT, 0.5)
    rh_list = _split(bottom - top - 4, [1] * len(rows), 4)
    ry = top + 4
    row_y = []
    for i, lab in enumerate(rows):
        sp = _rect(slide, x, ry, label_w, rh_list[i], fill=COLOR_PRIMARY)
        _put_text(sp, lab, size=size, color=COLOR_WHITE, bold=True, align="center", anchor="middle",
                  dense=True, where=f"gantt行[{i}]")
        _line(slide, gx, ry + rh_list[i] + 2, gx + gw, ry + rh_list[i] + 2, COLOR_GRAY_LIGHT, 0.5)
        row_y.append(ry)
        ry += rh_list[i] + 4
    lanes = {}
    for b in bars:
        lanes[b[0]] = max(lanes.get(b[0], 0), (b[5] if len(b) > 5 else 0) + 1)
    sty = {"primary": (COLOR_PRIMARY_MID, COLOR_WHITE), "mid": (COLOR_PRIMARY, COLOR_WHITE),
           "pale": (COLOR_PRIMARY_PALE, COLOR_PRIMARY)}
    for b in bars:
        r, bs, be, lab = b[:4]
        style = b[4] if len(b) > 4 else "primary"
        lane = b[5] if len(b) > 5 else 0
        lh = rh_list[r] / lanes[r]
        bh = min(24, lh - 6)
        bx0, bx1 = dx(bs), dx(be + timedelta(days=1))
        fill, col = sty[style]
        sp = _rect(slide, bx0, row_y[r] + lane * lh + (lh - bh) / 2, bx1 - bx0, bh, fill=fill,
                   shape=MSO_SHAPE.PENTAGON)
        _put_text(sp, lab, size=11, color=col, bold=True, anchor="middle", dense=True, inset=(4, 0),
                  where=f"ganttバー「{lab}」（バーが短すぎる場合はラベルを短く）")
    for md, lab in milestones:
        mx = dx(md)
        _line(slide, mx, top, mx, bottom, COLOR_ACCENT, 1.0, dash=True)
        lx = min(max(mx - 80, x), x + w - 160)
        _textbox(slide, lx, bottom, 160, 18, f"▲{lab}", size=10, bold=True, color=COLOR_ACCENT,
                 align="center", dense=True, where="ganttマイルストーン")
    return row_y, rh_list


# ================================================================ 部品用の共通ユーティリティ
def _alpha(shape, pct):
    """塗りに透過率を付ける（pct=不透明度 0-100）"""
    srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(pct * 1000))})
    srgb.append(a)


def _box(slide, x, y, w, h, body, style="text", size=SZ_BODY, align=None, anchor="middle",
         where="box", shape=MSO_SHAPE.RECTANGLE, bullets=None, dense=False, inset=(6, 4), line_w=LINE_W):
    """_CELL スタイルの箱に文字を入れる"""
    fill, line, col, bold, al, bul = _CELL[style]
    sp = _rect(slide, x, y, w, h, fill=fill, line=line, shape=shape, line_w=line_w)
    is_list = isinstance(body, list)
    if body is not None and body != "":
        _put_text(sp, body, size=size, color=col, bold=bold, align=align or ("left" if is_list else al),
                  anchor=anchor, bullets=(bul and is_list) if bullets is None else bullets, where=where,
                  dense=dense, inset=inset)
    return sp


def _circle(slide, x, y, d, text=None, fill=COLOR_PRIMARY, color=COLOR_WHITE, size=14, line=None, bold=True):
    sp = _rect(slide, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL)
    if text is not None:
        _put_text(sp, str(text), size=size, color=color, bold=bold, align="center", anchor="middle",
                  inset=(0, 0), dense=True, where="circle", check=False)
    return sp


def _arrow_line(slide, x1, y1, x2, y2, color=COLOR_GRAY_DARK, width=1.25, head=True, dash=False,
                elbow=False, both=False):
    """矢印付きの線（elbow=True でカギ線）"""
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    ln = slide.shapes.add_connector(kind, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    ln.line.color.rgb = _rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    lnx = ln.line._get_or_add_ln()
    if both:
        lnx.append(lnx.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if head:
        lnx.append(lnx.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    _no_effect(ln)
    ln.name = "arrow"
    return ln


def _poly(slide, pts, fill=COLOR_PRIMARY, line=None):
    fb = slide.shapes.build_freeform(Pt(pts[0][0]), Pt(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Pt(px), Pt(py)) for px, py in pts[1:]], close=True)
    sp = fb.convert_to_shape()
    sp.shadow.inherit = False
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = _rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = _rgb(line)
        sp.line.width = Pt(LINE_W)
    else:
        sp.line.fill.background()
    return sp


def _label(slide, x, y, w, h, body, size=SZ_BODY, bold=False, color=COLOR_TEXT, align="left",
           anchor="top", bullets=False, dense=False, where="label"):
    return _textbox(slide, x, y, w, h, body, size=size, bold=bold, color=color, align=align,
                    anchor=anchor, bullets=bullets, dense=dense, where=where)


# ================================================================ 検査（build_deck から呼ぶ）
def check_presentation(prs):
    """デッキ全体のルール違反を検出: (errors, warnings)"""
    errors, warns = [], []
    for si, slide in enumerate(prs.slides, start=1):
        boxes = []
        for sp in slide.shapes:
            l, t = sp.left / 12700, sp.top / 12700
            r, b = l + sp.width / 12700, t + sp.height / 12700
            if l < -0.5 or t < -0.5 or r > SLIDE_W + 0.5 or b > SLIDE_H + 0.5:
                errors.append(f"slide{si}: 「{sp.name}」がスライド外にはみ出しています ({l:.0f},{t:.0f})-({r:.0f},{b:.0f})")
            chrome = sp.name.startswith("chrome:")
            if sp.has_text_frame and sp.text_frame.text.strip():
                if not chrome and sp.shape_type != 1 and not sp.name.startswith(("arrow", "deco")):
                    boxes.append((sp.name, l, t, r, b, sp.text_frame.text[:15]))
                for p in sp.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size is not None and not chrome and run.font.size.pt < MIN_DENSE_PT:
                            errors.append(f"slide{si}: フォント {run.font.size.pt}pt < {MIN_DENSE_PT}pt 「{run.text[:15]}」")
                        try:
                            c = str(run.font.color.rgb) if run.font.color and run.font.color.type else None
                        except AttributeError:
                            c = None
                        if c and c.upper() not in PALETTE:
                            warns.append(f"slide{si}: パレット外の文字色 #{c} 「{run.text[:15]}」")
            if sp.shape_type == 1 or sp.has_text_frame:
                try:
                    if sp.fill.type == 1:
                        c = str(sp.fill.fore_color.rgb).upper()
                        if c not in PALETTE:
                            warns.append(f"slide{si}: パレット外の塗り #{c}「{sp.name}」")
                except Exception:
                    pass
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, c = boxes[i], boxes[j]
                ow = min(a[3], c[3]) - max(a[1], c[1])
                oh = min(a[4], c[4]) - max(a[2], c[2])
                if ow > 2 and oh > 2:
                    warns.append(f"slide{si}: テキスト同士が重なっています「{a[5]}」と「{c[5]}」")
    return errors, warns


# ================================================================ 追加部品（カテゴリ別モジュール）
from parts_text import *       # noqa: E402,F401,F403
from parts_table import *      # noqa: E402,F401,F403
from parts_flow import *       # noqa: E402,F401,F403
from parts_structure import *  # noqa: E402,F401,F403
from parts_data import *       # noqa: E402,F401,F403
from parts_media import *      # noqa: E402,F401,F403
from parts_connect import *    # noqa: E402,F401,F403
from parts_diagram import *    # noqa: E402,F401,F403
from parts_biz import *        # noqa: E402,F401,F403
