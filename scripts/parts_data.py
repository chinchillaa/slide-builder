# -*- coding: utf-8 -*-
"""parts_data - 数値・グラフの部品（KPI・比較・進捗・ウォーターフォール・散布図など）"""
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from pptx_helpers import *  # noqa: F401,F403
from pptx_helpers import (fit_one_line, _box, _check_size, _line, _put_text, _rect, _rgb, _set_ea, _textbox, _triangle,
                          chart, kpi, part, text_width)
from parts_text import split_cols, split_rows, titled_box


def _style_axes(ch, size, number_format="#,##0"):
    ch.font.name, ch.font.size = FONT, Pt(size)
    for rpr in ch._chartSpace.iter(qn("a:defRPr")):
        _set_ea(rpr)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb(COLOR_GRAY_LIGHT)
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(size)
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False


@part("数値", "KPI カードを横に並べる")
def kpi_row(slide, area, items, unit=""):
    """items: [(値, ラベル, 注記 or None)]"""
    for (v, lab, note), a in zip(items, split_cols(area, [1] * len(items))):
        kpi(slide, a, v, lab, note=note, unit=unit)


@part("数値", "前後比較の数値（Before → After と増減）")
def stat_compare(slide, area, before, after, delta=None, unit="", size=36):
    """before/after = (ラベル, 値)。delta は "▲40%" などの表示文字列"""
    x, y, w, h = area
    a, b = split_cols(area, [1, 1], 60)
    for (lab, val), (cx, cy, cw, ch), col in ((before, a, COLOR_TEXT_SUB), (after, b, COLOR_PRIMARY)):
        _rect(slide, cx, cy, cw, ch, fill=COLOR_WHITE, line=COLOR_GRAY)
        _textbox(slide, cx, cy + 8, cw, 24, lab, size=14, bold=True, align="center", where="stat_compareラベル")
        _textbox(slide, cx, cy + 34, cw, ch - 42, f"{val}{unit}", size=fit_one_line(f"{val}{unit}", cw - 10, size), bold=True, color=col,
                 align="center", anchor="middle", where="stat_compare値")
    _triangle(slide, a[0] + a[2] + 18, y + h / 2 - 22, 24, 44, fill=COLOR_PRIMARY_MID)
    if delta:
        tw = text_width(delta, 14) + 20
        sp = _rect(slide, b[0] + b[2] - tw - 8, b[1] + 8, tw, 24, fill=COLOR_WHITE, line=COLOR_ACCENT,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        _put_text(sp, delta, size=14, bold=True, color=COLOR_ACCENT, align="center", anchor="middle",
                  inset=(2, 0), where="stat_compare増減")


@part("数値", "進捗バー（％）と目標線")
def progress_bars(slide, area, items, target=None, label_ratio=0.3, size=12):
    """items: [(ラベル, 0-100)]  target: 目標値(%)"""
    x, y, w, h = area
    lw = w * label_ratio
    bw = w - lw - 60
    rows = split_rows(area, [1] * len(items), 6)
    for i, ((lab, pct), (rx, ry, rw, rh)) in enumerate(zip(items, rows)):
        bh = min(rh * 0.55, 22)
        by = ry + (rh - bh) / 2
        _textbox(slide, rx, ry, lw - 8, rh, lab, size=size, anchor="middle", where=f"progress_bars[{i}]")
        _rect(slide, rx + lw, by, bw, bh, fill=COLOR_GRAY_PALE)
        if pct > 0:
            _rect(slide, rx + lw, by, bw * min(pct, 100) / 100, bh, fill=COLOR_PRIMARY_MID)
        _textbox(slide, rx + lw + bw + 6, ry, 54, rh, f"{pct:g}%", size=size, bold=True, color=COLOR_PRIMARY,
                 anchor="middle", where="progress_bars値")
    if target is not None:
        tx = x + lw + bw * target / 100
        _line(slide, tx, y, tx, y + h, COLOR_ACCENT, 1.25, dash=True)
        _textbox(slide, tx - 40, y - 18, 80, 18, f"目標 {target:g}%", size=10, bold=True, color=COLOR_ACCENT,
                 align="center", dense=True, where="progress_bars目標")


@part("数値", "ランキングの横棒（図形で描く。強調行あり）")
def rank_bars(slide, area, items, unit="", highlight=None, fmt="{:,.0f}", label_ratio=0.25, size=12):
    """items: [(ラベル, 値)] 上から順に表示"""
    x, y, w, h = area
    vmax = max(v for _, v in items) or 1
    lw = w * label_ratio
    bw = w - lw - 70
    for i, ((lab, v), (rx, ry, rw, rh)) in enumerate(zip(items, split_rows(area, [1] * len(items), 4))):
        on = highlight == i
        bh = min(rh * 0.65, 24)
        _textbox(slide, rx, ry, lw - 8, rh, f"{i + 1}. {lab}", size=size, bold=on, anchor="middle",
                 color=COLOR_PRIMARY if on else COLOR_TEXT, where=f"rank_bars[{i}]")
        _rect(slide, rx + lw, ry + (rh - bh) / 2, max(bw * v / vmax, 1), bh,
              fill=COLOR_PRIMARY if on else COLOR_PRIMARY_LIGHT)
        _textbox(slide, rx + lw + bw * v / vmax + 6, ry, 70, rh, fmt.format(v) + unit, size=size, bold=on,
                 anchor="middle", dense=True, where="rank_bars値")


@part("数値", "100% 帯グラフ（構成比を1本の帯で）")
def percent_bar(slide, area, segments, size=12):
    """segments: [(ラベル, 割合%)]"""
    x, y, w, h = area
    bh = min(44, h * 0.4)
    by = y + (h - bh) / 2
    total = sum(p for _, p in segments) or 1
    cols = [COLOR_PRIMARY, COLOR_PRIMARY_MID, COLOR_PRIMARY_LIGHT, COLOR_GRAY, COLOR_GRAY_LIGHT]
    cx = x
    for i, (lab, p) in enumerate(segments):
        sw = w * p / total
        col = cols[i % len(cols)]
        sp = _rect(slide, cx, by, sw, bh, fill=col, line=COLOR_WHITE)
        if sw > 34:
            _put_text(sp, f"{p:g}%", size=size, bold=True, align="center", anchor="middle", inset=(1, 0),
                      color=COLOR_WHITE if i % len(cols) < 2 else COLOR_TEXT, dense=True, where="percent_bar値")
        up = i % 2 == 0
        ly = by - 26 if up else by + bh + 4
        _textbox(slide, cx + sw / 2 - 70, ly, 140, 22, lab, size=size, align="center", dense=True,
                 anchor="bottom" if up else "top", where=f"percent_bar[{i}]")
        cx += sw


@part("グラフ", "ウォーターフォール（増減の積み上げ。ネイティブ）")
def waterfall(slide, area, start, steps, end_label="合計", unit=None, title=None, size=12, number_format="#,##0"):
    """start=(ラベル, 値)  steps=[(ラベル, 増減)]"""
    x, y, w, h = area
    if title or unit:
        if title:
            _textbox(slide, x, y, w, 24, title, size=14, bold=True, align="center", where="グラフタイトル")
        if unit:
            _textbox(slide, x + w - 160, y + 22, 160, 16, f"（単位：{unit}）", size=10, align="right",
                     color=COLOR_TEXT_SUB, dense=True, where="グラフ単位")
        y, h = y + 40, h - 40
    cats = [start[0]] + [s[0] for s in steps] + [end_label]
    base, tot, up, down = [0], [start[1]], [0], [0]
    run = start[1]
    for _, d in steps:
        if d >= 0:
            base.append(run); up.append(d); down.append(0)
        else:
            base.append(run + d); up.append(0); down.append(-d)
        tot.append(0)
        run += d
    base.append(0); tot.append(run); up.append(0); down.append(0)
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in (("base", base), ("合計", tot), ("増加", up), ("減少", down)):
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    _style_axes(ch, size, number_format)
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 50, 100
    for s, col in zip(plot.series, (None, COLOR_PRIMARY, COLOR_PRIMARY_MID, COLOR_GRAY)):
        if col is None:
            s.format.fill.background()
            continue
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = _rgb(col)
        dl = s.data_labels
        dl.show_value = True
        dl.number_format = number_format + ";-" + number_format + ";;"
        dl.number_format_is_linked = False
        dl.font.size = Pt(size)
        dl.font.bold = True
        dl.font.color.rgb = _rgb(COLOR_WHITE if col != COLOR_GRAY else COLOR_TEXT)
        dl.position = XL_LABEL_POSITION.CENTER
    ch.category_axis.tick_labels.font.size = Pt(size)
    ch.category_axis.format.line.color.rgb = _rgb(COLOR_GRAY_DARK)
    return gf


def _xy_style(ch, size, x_title, y_title):
    _style_axes(ch, size)
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(size)
    ca.format.line.color.rgb = _rgb(COLOR_GRAY_DARK)
    for ax, t in ((ca, x_title), (ch.value_axis, y_title)):
        if t:
            ax.has_title = True
            ax.axis_title.text_frame.text = t
            r = ax.axis_title.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(size)
            r.font.bold = False
            r.font.name = FONT
            _set_ea(r.font._rPr)


@part("グラフ", "散布図（ネイティブ）")
def scatter_chart(slide, area, series, x_title=None, y_title=None, size=12):
    """series: {"系列名": [(x, y), ...]}"""
    x, y, w, h = area
    _check_size(size, True, "scatter_chart")
    cd = XyChartData()
    for name, pts in series.items():
        s = cd.add_series(name)
        for px, py in pts:
            s.add_data_point(px, py)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    _xy_style(ch, size, x_title, y_title)
    for i, s in enumerate(ch.plots[0].series):
        col = SERIES_COLORS[i % len(SERIES_COLORS)]
        s.format.line.fill.background()
        s.marker.size = 9
        s.marker.format.fill.solid()
        s.marker.format.fill.fore_color.rgb = _rgb(col)
        s.marker.format.line.color.rgb = _rgb(COLOR_WHITE)
    return gf


@part("グラフ", "バブルチャート（ネイティブ。3つ目の値を円の大きさで）")
def bubble_chart(slide, area, series, x_title=None, y_title=None, size=12):
    """series: {"系列名": [(x, y, 大きさ), ...]}"""
    x, y, w, h = area
    _check_size(size, True, "bubble_chart")
    cd = BubbleChartData()
    for name, pts in series.items():
        s = cd.add_series(name)
        for px, py, pz in pts:
            s.add_data_point(px, py, pz)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE, Pt(x), Pt(y), Pt(w), Pt(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = len(series) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    _xy_style(ch, size, x_title, y_title)
    for i, s in enumerate(ch.plots[0].series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = _rgb(SERIES_COLORS[i % len(SERIES_COLORS)])
    return gf


@part("グラフ", "グラフ＋右に示唆ボックス")
def chart_with_insight(slide, area, kind, categories, series, insights, insight_title="示唆", ratio=0.62, **kw):
    """chart() の引数に加えて insights（箇条書き）を受け取る"""
    a, b = split_cols(area, [ratio, 1 - ratio], 20)
    chart(slide, a, kind, categories, series, **kw)
    titled_box(slide, b, insight_title, insights, style="primary", size=14)


@part("数値", "大きな数字＋説明文（インパクトのある1つの数値）")
def big_number(slide, area, number, unit="", desc="", label=None, size=64):
    x, y, w, h = area
    nw = w * 0.45
    size = fit_one_line(f"{number}{unit}", nw - 8, min(size, (h - (26 if label else 0) - 8) / 1.3))   # 高さに合わせて数字の大きさを決める（本文の下限は別途検査）
    if label:
        _textbox(slide, x, y, nw, 26, label, size=14, bold=True, color=COLOR_TEXT_SUB, align="center",
                 where="big_numberラベル")
    _textbox(slide, x, y + (26 if label else 0), nw, h - (26 if label else 0),
             f"{number}" + (f"**{unit}**" if unit else ""), size=size, bold=True, color=COLOR_PRIMARY,
             align="center", anchor="middle", where="big_number数値")
    _line(slide, x + nw + 10, y + h * 0.15, x + nw + 10, y + h * 0.85, COLOR_GRAY, 1)
    _textbox(slide, x + nw + 26, y, w - nw - 26, h, desc, size=16, anchor="middle",
             bullets=isinstance(desc, list), where="big_number説明")


@part("数値", "ドーナツ型の達成率（中央に％）")
def donut_kpi(slide, area, pct, label, size=12):
    x, y, w, h = area
    d = min(w, h - 28)
    cx = x + (w - d) / 2
    cd = CategoryChartData()
    cd.categories = ["達成", "残り"]
    cd.add_series("率", [pct, max(100 - pct, 0)])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Pt(cx), Pt(y), Pt(d), Pt(d), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    pts = ch.plots[0].series[0].points
    for i, col in enumerate((COLOR_PRIMARY_MID, COLOR_GRAY_PALE)):
        pts[i].format.fill.solid()
        pts[i].format.fill.fore_color.rgb = _rgb(col)
        pts[i].format.line.fill.background()
    hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "70")
    _textbox(slide, cx + d * 0.2, y + d * 0.3, d * 0.6, d * 0.4, f"{pct:g}%", size=max(14, min(28, d * 0.16)),
             bold=True, color=COLOR_PRIMARY, align="center", anchor="middle", where="donut_kpi値")
    _textbox(slide, x, y + d + 2, w, 24, label, size=size + 1, bold=True, align="center", where="donut_kpiラベル")


@part("数値", "推移線付きの KPI カード（小さな折れ線）")
def kpi_trend(slide, area, value, label, values, unit="", note=None):
    """values: 推移の数値リスト（古い順）"""
    x, y, w, h = area
    _rect(slide, x, y, w, h, fill=COLOR_WHITE, line=COLOR_GRAY_DARK)
    _textbox(slide, x + 10, y + 6, w - 20, 24, label, size=14, bold=True, where="kpi_trendラベル")
    vw = w * 0.45
    vs = fit_one_line(f"{value}{unit}", vw - 8, min(36, (h - 40 - (24 if note else 0)) / 1.3))
    _textbox(slide, x + 10, y + 32, vw, h - 40 - (24 if note else 0), f"{value}{unit}", size=vs, bold=True,
             color=COLOR_PRIMARY, anchor="middle", where="kpi_trend値")
    if note:
        _textbox(slide, x + 10, y + h - 28, w - 20, 22, note, size=12, color=COLOR_TEXT_SUB, where="kpi_trend注記")
    cd = CategoryChartData()
    cd.categories = [str(i) for i in range(len(values))]
    cd.add_series(label, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Pt(x + vw + 14), Pt(y + 34), Pt(w - vw - 24),
                                Pt(h - 44 - (24 if note else 0)), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.visible = False
    s = ch.plots[0].series[0]
    s.format.line.color.rgb = _rgb(COLOR_PRIMARY_MID)
    s.format.line.width = Pt(2.25)
    s.smooth = False
    return gf
